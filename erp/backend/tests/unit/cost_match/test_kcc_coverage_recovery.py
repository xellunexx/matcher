from __future__ import annotations

import csv
import json
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy.dialects import postgresql

from app.modules.cost_match.matcher import Candidate, _price_day, best_match, score_match, work_rate_factor
from app.modules.cost_match.retrieval_bg import literal_description, retrieval_plan, sql_predicate
from app.modules.cost_match.work_catalog import canonical_work
from app.modules.costs.models import CostItem
from scripts.kcc_match_bg import run


def quotation(code, text, rate, *, unit='м', source='operator_pricelist', **extra):
    operator = source in {'operator_pricelist', 'estimate_confirmed'}
    return Candidate(code, text, unit, {
        'unit_rate': str(rate), 'currency': 'EUR', 'source': source,
        'auto_pricing_eligible': operator,
        'provenance_status': 'approved' if operator else 'pending_review', **extra,
    })


BATTERY = 'Доставка и монтаж на акумолаторна батерия:\n12V 65Ah\nСертифицирана EN54'


@pytest.mark.parametrize('line_break', ['\n', '\r\n', '\r'])
def test_literal_unknown_equipment_retrieval_normalizes_only_line_endings(line_break):
    quote = BATTERY.replace('\n', line_break)
    plan = retrieval_plan(BATTERY)
    assert plan.exact_text == literal_description(quote)
    assert plan.matches(quote)
    result = best_match(BATTERY, [quotation('battery', quote, 124, unit='бр.')], query_unit='бр.')
    assert result.is_confident
    assert result.candidate.ref == 'battery'


@pytest.mark.parametrize(('old', 'new'), [('65Ah', '7Ah'), ('12V', '24V'), ('EN54', 'EN55')])
def test_unknown_equipment_literal_retrieval_keeps_all_technical_requirements(old, new):
    assert not retrieval_plan(BATTERY).matches(BATTERY.replace(old, new).replace('\n', '\r\n'))


def test_literal_sql_uses_the_same_line_ending_normalization():
    plan = retrieval_plan(BATTERY)
    statement = sql_predicate(plan, CostItem).compile(dialect=postgresql.dialect())
    assert str(statement).count('replace(') == 2
    assert BATTERY.lower() in statement.params.values()
    assert '\r\n' in statement.params.values()
    assert '\r' in statement.params.values()


@pytest.mark.parametrize('section', ['2,5', '16'])
def test_cable_connection_binds_equivalent_destination_and_section(section):
    query = f'Свързване проводник към табло до {section}мм²'
    quote = f'Свързване проводник към съоръжение до {section}мм²'
    work = canonical_work(query)
    assert work.attributes['connection_target'] == ('electrical_equipment',)
    assert work.specs['cable_section_max_mm2'] == (section.replace(',', '.'),)
    assert retrieval_plan(query).matches(quote)
    result = best_match(query, [quotation('connection', quote, 1, unit='бр.')], query_unit='бр.')
    assert result.is_confident
    assert result.pool_refs == {'connection'}


@pytest.mark.parametrize('quote', [
    'Свързване проводник към съоръжение до 16мм²',
    'Свързване проводник към тръба до 2,5мм²',
    'Свързване проводник към съоръжение до 2,5мм² с ухо',
    'Свързване проводник към съоръжение до 2,5мм² без демонтаж',
])
def test_cable_connection_rejects_size_target_accessory_and_eligibility_changes(quote):
    score = score_match('Свързване проводник към табло до 2,5мм²', quote,
                        query_unit='бр.', candidate_unit='бр.')
    assert score.confidence < 0.85


GRANITE = 'Полагане на гранит -стълби - включително грундиране и замазка'


@pytest.mark.parametrize('place', ['(топла връзка)', '(експозиционна зала)', 'към сутерен'])
def test_granite_stair_destination_is_audited_as_room_context(place):
    query = GRANITE.replace(' - включително', f' {place} - включително')
    work = canonical_work(query)
    assert work.contexts['room']
    assert 'basement' not in work.locations
    result = best_match(query, [quotation('granite', GRANITE, 44, unit='м²')], query_unit='м²')
    assert result.is_confident
    assert result.pool_refs == {'granite'}


@pytest.mark.parametrize('quote', [
    'Полагане на гранит -стълби по външна фасада - включително грундиране и замазка',
    'Полагане на гранитогрес -стълби - включително грундиране и замазка',
    'Полагане на гранит -стълби',
])
def test_granite_stair_context_does_not_drop_surface_material_or_bundled_work(quote):
    result = best_match(GRANITE, [quotation('wrong', quote, 44, unit='м²')], query_unit='м²')
    assert not result.is_confident


PIPE = 'Тръба PE PN10; Ф63'
PIPE_QUOTE = 'ПЕВП Тръба Ф63 SDR17 / PN10 - 100m-ролка'


def test_incomplete_pipe_rival_cannot_impeach_a_complete_refinement():
    candidates = [quotation('pipe1', PIPE_QUOTE, 14.5), quotation('pipe2', PIPE_QUOTE, 14.5),
                  quotation('unknown-pn', 'Заваряване на ПЕВП тръба Ф63 SDR11', 8)]
    result = best_match(PIPE, candidates, query_unit='м')
    assert result.is_confident
    assert not result.pool_diverged
    assert result.pool_refs == {'pipe1', 'pipe2'}
    assert result.median_rate == Decimal('14.5')


def test_complete_pipe_refinement_competitor_still_requires_review():
    result = best_match(PIPE, [quotation('sdr17', PIPE_QUOTE, 14.5),
                              quotation('sdr11', PIPE_QUOTE.replace('SDR17', 'SDR11'), 14.5)],
                        query_unit='м')
    assert not result.is_confident
    assert 'pool_divergence' in result.score.reasons


def test_refined_pipe_cohort_rejects_wrong_required_pressure_and_reference_authority():
    candidates = [quotation('pipe1', PIPE_QUOTE, 14.5), quotation('pipe2', PIPE_QUOTE, 14.5),
                  quotation('pn16', PIPE_QUOTE.replace('PN10', 'PN16'), 1),
                  quotation('reference', PIPE_QUOTE, 1, source='reference_web')]
    result = best_match(PIPE, candidates, query_unit='м')
    assert not result.is_confident
    assert result.pool_diverged
    assert result.pool_refs == {'pipe1', 'pipe2', 'reference'}
    assert any(candidate.ref == 'reference' for candidate, _ in result.scored_all)


def test_mixed_reference_evidence_is_preserved_for_application_review():
    text = 'Боядисване с латекс едноцветно две ръце'
    result = best_match(text, [quotation('ref', text, 6, unit='м²', source='reference_web'),
                              quotation('operator', text, 5, unit='м²')], query_unit='м²')
    assert not result.is_confident
    assert result.candidate.ref == 'operator'
    assert result.pool_refs == {'operator', 'ref'}
    assert result.pool_min == Decimal('5') and result.pool_max == Decimal('6')
    assert result.median_rate == Decimal('5.5')
    assert any(candidate.ref == 'ref' for candidate, _ in result.alternatives)


def test_ramp_slope_is_required_and_separates_price_cohorts():
    query = 'Новопроектирани бетонни рампи 5%'
    result = best_match(query, [quotation('five', query, 60, unit='м²'),
                               quotation('six', query.replace('5%', '6%'), 85, unit='м²')], query_unit='м²')
    assert result.is_confident
    assert result.pool_refs == {'five'}
    assert canonical_work(query).specs['cable_slope_percent'] == ('5',)


def test_same_slope_operator_disagreement_is_not_hidden():
    text = 'Новопроектирани бетонни рампи 6%'
    result = best_match(text, [quotation('a', text, 25, unit='м²'),
                              quotation('b', text, 85, unit='м²')], query_unit='м²')
    assert not result.is_confident
    assert result.pool_diverged


def test_import_date_is_not_price_authority_but_live_confirmation_date_is():
    assert _price_day(quotation('import', PIPE_QUOTE, 10, created_at='2026-10-08T12:00:00')) == ''
    assert _price_day(quotation('vintage', PIPE_QUOTE, 10, price_as_of='2025-01-02',
                                created_at='2026-10-08T12:00:00')) == '2025-01-02'
    assert _price_day(quotation('ruling', PIPE_QUOTE, 10, source='estimate_confirmed',
                                created_at='2026-10-08T12:00:00')) == '2026-10-08'


WALL = 'Разваляне на 25 см. тухлен зид'
WALL_QUOTE = 'Разваляне на тухлена зидария 1 тухла. крайна цена'


@pytest.mark.parametrize(('source_unit', 'target_unit', 'expected'), [
    ('м³', 'м²', '0.25'), ('100 м³', 'м²', '0.0025'), ('м²', 'м³', '4'),
])
def test_equivalent_explicit_brick_wall_geometry_is_not_factor_one(source_unit, target_unit, expected):
    factor = work_rate_factor(source_unit, target_unit, query_text=WALL, candidate_text=WALL_QUOTE)
    assert factor == Decimal(expected)
    assert canonical_work(WALL).attributes['wall_thickness_cm'] == ('25',)
    assert 'dim_len' not in canonical_work(WALL).specs


@pytest.mark.parametrize('quote', [
    'Разваляне на тухлена зидария 1/2 тухла',
    'Разваляне на тухлена зидария',
    'Очукване на издадени части от тухлена зидария 1 тухла',
    'Разваляне на тухлена зидария 1 тухла без демонтаж',
    'Направа на тухлена зидария 1 тухла',
    'Разваляне на бетонен зид с дебелина 25см',
])
def test_geometric_conversion_rejects_missing_or_wrong_work_and_dimensions(quote):
    assert work_rate_factor('м³', 'м²', query_text=WALL, candidate_text=quote) is None


def test_wall_height_is_not_evidence_of_wall_thickness():
    assert work_rate_factor('м³', 'м²', query_text='Разваляне на тухлен зид с височина 25см',
                            candidate_text=WALL_QUOTE) is None


def test_masonry_price_pool_converts_an_original_volume_price_honestly():
    result = best_match(WALL, [quotation('volume', WALL_QUOTE, 35.91, unit='м³')], query_unit='м²')
    assert result.is_confident
    assert result.pool_min == result.pool_max == Decimal('8.9775')
    assert result.candidate.payload['unit_rate'] == '35.91'
    assert result.candidate.unit == 'м³'


@pytest.mark.parametrize('query', [
    'Кофраж за плоча, греди и пояси', 'Кофраж за пояси',
])
def test_ring_beam_is_not_silently_priced_by_a_partial_slab_and_beam_quote(query):
    quote = 'Кофраж за ст.б. плоча и греди'
    result = best_match(query, [quotation('partial', quote, 50, unit='м²')], query_unit='м²')
    assert not result.is_confident
    assert 'ring_beam' in canonical_work(query).specs['cable_element']


def test_bare_pipe_description_does_not_authorize_a_supply_and_installation_bundle():
    result = best_match("Поцинковани тръби-Ф2''",
                        [quotation('bundle', 'Доставка и монтаж на поцинковани тръби 2"', 30.42)], query_unit='м')
    assert not result.is_confident
    assert 'work_scope_unknown' in result.score.reasons


@pytest.mark.parametrize(('source', 'high_rate', 'priced'), [
    ('reference_web', '1000', False), ('operator_pricelist', '1000', False),
    ('operator_pricelist', '5', True),
])
def test_replay_source_separation_and_import_conflicts(tmp_path, source, high_rate, priced):
    text = 'Боядисване с латекс едноцветно две ръце'
    corpus = tmp_path / 'corpus.csv'
    with corpus.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=['code', 'description', 'rate', 'unit',
                                                    'currency', 'source', 'created_at', 'price_as_of'])
        writer.writeheader()
        for code, rate, origin, day in [('op', '5', 'operator_pricelist', '2026-10-07'),
                                       ('other', high_rate, source, '2026-10-08')]:
            writer.writerow(dict(code=code, description=text, rate=rate, unit='м²', currency='EUR',
                                 source=origin, created_at=day, price_as_of=''))
    original, output = tmp_path / 'input.xlsx', tmp_path / 'output.xlsx'
    book = Workbook()
    book.active.append(['№', 'Описание', 'Ед. мярка', 'Количество'])
    book.active.append([1, text, 'м²', 10])
    book.save(original)
    counts = run(corpus, original, output, 1)
    result = load_workbook(output)
    assert result.active['E2'].value == (5 if priced else None)
    assert result.active['F2'].value == (50 if priced else None)
    assert counts['exact' if priced else 'review'] == 1
    audit = result['BG_MATCH_AUDIT']
    pool = json.loads(audit.cell(2, 23).value)
    expected_sources = {'operator_pricelist', 'reference_web'} if source == 'reference_web' else {'operator_pricelist'}
    assert {item['source'] for item in pool['observations']} == expected_sources
