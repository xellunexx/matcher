import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scripts.kcc_match_bg import run


@pytest.mark.parametrize(('rates', 'expected_rate', 'decision'), [
    ([10, '12.3'], 11.15, 'exact'),
    ([100, 10, '10.1', '10.2', '10.3', '10.4'], None, 'review'),
])
def test_replay_uses_previous_price_policy_without_a_20_percent_cohort(tmp_path, rates, expected_rate, decision):
    import json

    text = 'Демонтаж на осветителни тела'
    corpus = tmp_path / 'prices.csv'
    with corpus.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=['code', 'description', 'unit', 'rate', 'currency',
                                                   'source', 'classification'])
        writer.writeheader()
        for index, rate in enumerate(rates):
            writer.writerow(dict(code=f'quote-{index}', description=text, unit='бр.', rate=rate, currency='EUR',
                                 source='operator_pricelist',
                                 classification=json.dumps({'evidence_origin': f'project-{index}'})))
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    workbook = Workbook()
    workbook.active.append(['№', 'Описание', 'Ед. мярка', 'Количество'])
    workbook.active.append([1, text, 'бр.', 2])
    workbook.save(source)
    counts = run(corpus, source, output, 1)
    result = load_workbook(output, data_only=True)
    assert counts[decision] == 1
    assert result.active['E2'].value == expected_rate
    assert result.active['F2'].value == (None if expected_rate is None else expected_rate * 2)
    audit = result['BG_MATCH_AUDIT']
    column = [cell.value for cell in audit[1]].index('Equivalent price pool') + 1
    evidence = json.loads(audit.cell(2, column).value)
    assert 'cohort' not in evidence
    assert len(evidence['observations']) == len(rates)
    assert evidence['diverged'] == (decision == 'review')


def test_replay_prices_evidence_not_headers_or_partial_work(tmp_path: Path):
    corpus = tmp_path / 'prices.csv'
    with corpus.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=['code', 'description', 'unit', 'rate', 'currency', 'source', 'is_active'])
        writer.writeheader()
        for code, text, unit, rate in [
            ('OPR-brick', 'Иззиждане на отвори с тухла', 'м³', '8'),
            ('OPR-duct', 'Доставка и монтаж на въздуховод ф160', 'м', '20'),
            ('OPR-header', 'СМР-РАЗПРЕДЕЛЕНИЕ КОТА +3.60', 'м2', '999'),
            ('OPR-rampA', 'Бетонна рампа', 'м²', '85'),
            ('OPR-rampB', 'Бетонна рампа', 'кв.м.', '25'),
        ]:
            writer.writerow(dict(code=code, description=text, unit=unit, rate=rate,
                                 currency='EUR', source='operator_pricelist', is_active='True'))
    input_path, output_path = tmp_path / 'input.xlsx', tmp_path / 'output.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['№', 'Описание', 'Ед. мярка', 'Количество'])
    sheet.append([1, 'Иззиждане на отвори с тухла', 'м³', 6])
    sheet.append([2, 'Доставка и монтаж на въздуховод ф160 вкл. фасонни части', 'м', 10])
    sheet.append([3, 'СМР-РАЗПРЕДЕЛЕНИЕ КОТА +3.60', 'м2', 1])
    sheet.append([4, 'Бетонна рампа', 'м²', 1])
    sheet.append([5, 'Ръчно зададена цена', 'бр', 1, 12])
    workbook.save(input_path)
    counts = run(corpus, input_path, output_path, 100)
    priced = load_workbook(output_path, data_only=True)
    assert priced.active['E2'].value == 8
    assert priced.active['F2'].value == 48
    assert priced.active['E3'].value is None
    assert priced.active['E4'].value is None
    assert priced.active['E5'].value is None
    assert priced.active['E6'].value == 12
    assert counts['exact'] == 1
    assert counts['review'] == 2
    assert counts['structural_or_nonwork'] == 1
    assert 'BG_MATCH_AUDIT' in priced.sheetnames
    assert load_workbook(input_path, data_only=True).active['E2'].value is None


def test_replay_refuses_to_overwrite_input(tmp_path):
    source = tmp_path / 'book.xlsx'
    with pytest.raises(ValueError, match='overwrite'):
        run(tmp_path / 'missing.csv', source, source, 100)


@pytest.mark.parametrize('sources', [
    ['reference_web'], ['supplier_web_anchor'], ['operator_pricelist', 'reference_web'],
])
def test_replay_keeps_reference_only_price_pools_in_review(tmp_path, sources):
    text = 'Боядисване с латекс едноцветно две ръце'
    corpus = tmp_path / 'prices.csv'
    with corpus.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=['code', 'description', 'unit', 'rate', 'currency', 'source'])
        writer.writeheader()
        for index, source in enumerate(sources):
            writer.writerow(dict(code=f'price-{index}', description=text, unit='м²', rate='5',
                                 currency='EUR', source=source))
    original, output = tmp_path / 'input.xlsx', tmp_path / 'output.xlsx'
    workbook = Workbook()
    workbook.active.append(['№', 'Описание', 'Ед. мярка', 'Количество'])
    workbook.active.append([1, text, 'м²', 10])
    workbook.save(original)
    counts = run(corpus, original, output, 1)
    result = load_workbook(output)
    assert result.active['E2'].value is None
    assert result.active['F2'].value is None
    assert counts['review'] == 1
    assert 'reference_price' in result['BG_MATCH_AUDIT']['K2'].value


def test_replay_does_not_let_a_generic_book_price_veto_a_confirmed_ruling(tmp_path):
    text = 'Циркулационна помпа'
    corpus = tmp_path / 'prices.csv'
    with corpus.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=['code', 'description', 'unit', 'rate', 'currency', 'source'])
        writer.writeheader()
        for code, rate, source in [('book', '1000', 'operator_pricelist'),
                                   ('confirmed', '1600', 'estimate_confirmed')]:
            writer.writerow(dict(code=code, description=text, unit='бр.', rate=rate,
                                 currency='EUR', source=source))
    original, output = tmp_path / 'input.xlsx', tmp_path / 'output.xlsx'
    workbook = Workbook()
    workbook.active.append(['№', 'Описание', 'Ед. мярка', 'Количество'])
    workbook.active.append([1, text, 'бр.', 2])
    workbook.save(original)
    counts = run(corpus, original, output, 1)
    result = load_workbook(output)
    assert counts['exact'] == 1
    assert result.active['E2'].value == 1600
    assert result.active['F2'].value == 3200


@pytest.mark.parametrize('cached', [True, False])
def test_replay_uses_cached_formula_quantity_without_evaluating_or_replacing_formula(tmp_path, cached):
    import xml.etree.ElementTree as ET
    from zipfile import ZipFile

    text = 'Доставка и монтаж на въздуховод ф160'
    corpus = tmp_path / 'prices.csv'
    with corpus.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=['code', 'description', 'unit', 'rate', 'currency', 'source'])
        writer.writeheader()
        writer.writerow(dict(code='OPR-duct', description=text, unit='м', rate='20',
                             currency='EUR', source='operator_pricelist'))
    original, source, output = (tmp_path / name for name in ['original.xlsx', 'input.xlsx', 'output.xlsx'])
    workbook = Workbook()
    workbook.active.append(['№', 'Описание', 'Ед. мярка', 'Количество'])
    workbook.active.append([1, text, 'м', '=2+3'])
    workbook.save(original)
    with ZipFile(original) as before, ZipFile(source, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry.filename)
            if cached and entry.filename == 'xl/worksheets/sheet1.xml':
                root = ET.fromstring(data)
                ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                cell = root.find('.//s:c[@r="D2"]/s:v', ns)
                assert cell is not None
                cell.text = '5'
                data = ET.tostring(root, encoding='utf-8')
            after.writestr(entry, data)
    counts = run(corpus, source, output, 1)
    result = load_workbook(output)
    assert result.active['D2'].value == '=2+3'
    assert counts['structural_or_nonwork'] == 0
    assert result.active['E2'].value == (20 if cached else None)
    assert result.active['F2'].value == (100 if cached else None)
    if not cached:
        assert counts['review'] == 1
        assert result['BG_MATCH_AUDIT']['K2'].value == 'quantity_unavailable'
