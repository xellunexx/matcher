# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Register exports and letters are written in a worker thread, not on the loop.

One worker serves every user of an install, and a file written on its event
loop answers nobody else until the last row is on the page. The whole active
cost database, up to 50 000 contacts and up to 25 000 clash results are the
largest of these; a tender letter is small but draws a full PDF. The queries
stay on the loop with their session, and only plain values cross into the
thread.

Each test swaps a function the writer calls for a spy that records the thread
it ran on, calls the route or service, and asserts that none of those calls ran
on the test's own thread, which is the loop's thread.
"""

from __future__ import annotations

import io
import threading
import uuid
from collections.abc import Callable
from decimal import Decimal
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock, MagicMock

import pytest


def _recording(seen: list[int], real: Callable[..., Any]) -> Callable[..., Any]:
    def spy(*args: Any, **kwargs: Any) -> Any:
        seen.append(threading.get_ident())
        return real(*args, **kwargs)

    return spy


def _session_returning(items: list[Any]) -> MagicMock:
    result = MagicMock()
    result.scalars.return_value.all.return_value = items
    session = MagicMock()
    session.execute = AsyncMock(return_value=result)
    return session


async def _body(response: Any) -> bytes:
    chunks = []
    async for chunk in response.body_iterator:
        chunks.append(chunk if isinstance(chunk, bytes) else chunk.encode("utf-8"))
    return b"".join(chunks)


def _cells(body: bytes) -> list[Any]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(body))
    return [cell for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for cell in row]


@pytest.mark.asyncio
async def test_the_cost_database_is_written_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.modules.costs import router

    item = SimpleNamespace(
        code="C-001",
        description="=HYPERLINK(1)",
        unit="m3",
        rate=Decimal("185.50"),
        currency="EUR",
        source="cwicr",
        region="DE",
    )
    appended: list[int] = []
    saved: list[int] = []
    monkeypatch.setattr(router, "_append_rows", _recording(appended, router._append_rows))
    monkeypatch.setattr(router, "_save_workbook", _recording(saved, router._save_workbook))

    response = await router.export_cost_database(session=_session_returning([item]), _user_id="u1")
    values = _cells(await _body(response))

    assert "C-001" in values
    assert "'=HYPERLINK(1)" in values
    assert 185.5 in values
    assert appended and saved, "the sheet was never written"
    assert threading.get_ident() not in appended + saved


@pytest.mark.asyncio
async def test_the_contacts_workbook_is_written_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.modules.contacts import router

    contact = SimpleNamespace(
        company_name="Riverside Concrete",
        first_name="Ana",
        last_name="Lopez",
        contact_type="subcontractor",
        primary_email="ana@example.com",
        primary_phone="+49 30 1234",
        country_code="DE",
        vat_number="DE123456789",
        prequalification_status="approved",
        payment_terms_days=30,
    )
    seen: list[int] = []
    monkeypatch.setattr(router, "_is_admin", AsyncMock(return_value=True))
    monkeypatch.setattr(router, "_render_contacts_xlsx", _recording(seen, router._render_contacts_xlsx))

    response = await router.export_contacts(session=_session_returning([contact]), user_id="u1", _perm=None)
    values = _cells(await _body(response))

    assert "Riverside Concrete" in values
    assert "Company" in values
    assert "DE123456789" in values
    assert seen, "the workbook writer never ran"
    assert threading.get_ident() not in seen


@pytest.mark.asyncio
async def test_the_clash_csv_is_written_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.modules.clash import router

    result = SimpleNamespace(
        a_name="Wall 12",
        a_discipline="ARC",
        b_name="=Duct 7",
        b_discipline="MEP",
        clash_type="hard",
        severity=None,
        penetration_m=0.12,
        distance_m=0.0,
        status="open",
        assigned_to=None,
        due_date=None,
    )
    service = SimpleNamespace(
        get_run=AsyncMock(return_value=SimpleNamespace(name="Level 1")),
        list_results=AsyncMock(return_value=([result], 1)),
    )
    seen: list[int] = []
    monkeypatch.setattr(router, "_require_project_access", AsyncMock(return_value=None))
    monkeypatch.setattr(router, "neutralise_formula", _recording(seen, router.neutralise_formula))

    response = await router.export_csv(
        project_id=uuid.uuid4(),
        run_id=uuid.uuid4(),
        user_id="u1",
        session=MagicMock(),
        service=service,
        status_filter=None,
        clash_type=None,
        severity=None,
    )
    lines = (await _body(response)).decode("utf-8").splitlines()

    assert lines[0].startswith("#,Element A,")
    assert lines[1] == "1,Wall 12,ARC,'=Duct 7,MEP,hard,medium,0.12,0.0,open,,"
    assert seen, "the CSV writer never ran"
    assert threading.get_ident() not in seen


@pytest.mark.asyncio
async def test_the_award_letter_is_drawn_off_the_loop(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.modules.tendering import pdf_documents
    from app.modules.tendering.service import TenderingService

    package_id = uuid.uuid4()
    service = TenderingService.__new__(TenderingService)
    service.get_package = AsyncMock(  # type: ignore[method-assign]
        return_value=SimpleNamespace(name="Shell and core", metadata_={})
    )
    service.get_bid = AsyncMock(  # type: ignore[method-assign]
        return_value=SimpleNamespace(
            package_id=package_id,
            company_name="Riverside Concrete",
            contact_email="bids@example.com",
            total_amount="125000.00",
            currency="EUR",
            notes=None,
        )
    )
    service._project_name_and_currency = AsyncMock(  # type: ignore[method-assign]
        return_value=("Tower A", "EUR")
    )
    seen: list[int] = []
    monkeypatch.setattr(
        pdf_documents,
        "generate_award_letter_pdf",
        _recording(seen, pdf_documents.generate_award_letter_pdf),
    )

    pdf, filename = await service.build_award_letter_pdf(package_id, uuid.uuid4())

    assert pdf.startswith(b"%PDF")
    assert filename.endswith(".pdf")
    assert seen, "the letter was never drawn"
    assert threading.get_ident() not in seen


@pytest.mark.asyncio
async def test_the_punch_list_pdf_gets_plain_copies_not_orm_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    from datetime import date

    from app.modules.punchlist import service as punch_service
    from app.modules.punchlist.models import PunchItem

    project_id = uuid.uuid4()
    item = PunchItem(
        project_id=project_id,
        title="Touch up paint at door 1.04",
        description="Scuffed frame",
        status="open",
        priority="high",
        category="finishes",
        trade="painting",
        assigned_to=None,
        due_date=date(2026, 10, 1),
        resolution_notes="Repainted",
        photos=["punch/door-104.jpg"],
        metadata_={"code": "P-17", "sheet_id": "A-101"},
        document_id=None,
        page=2,
        location_x=0.25,
        location_y=0.75,
        reopen_history=[{"reopened_at": "2026-09-20", "previous_status": "closed", "reopened_by": "site"}],
    )
    handed: list[list[Any]] = []
    seen: list[int] = []

    def _spy(real: Callable[..., Any], position: int) -> Callable[..., Any]:
        def spy(*args: Any, **kwargs: Any) -> Any:
            seen.append(threading.get_ident())
            handed.append(list(args[position]))
            return real(*args, **kwargs)

        return spy

    monkeypatch.setattr(punch_service, "_build_reportlab_pdf", _spy(punch_service._build_reportlab_pdf, 1))
    monkeypatch.setattr(punch_service, "_build_minimal_pdf", _recording(seen, punch_service._build_minimal_pdf))
    text_calls: list[list[Any]] = []
    real_text = punch_service._render_punchlist_text

    def text_spy(pid: Any, items: Any, names: Any) -> Any:
        text_calls.append(list(items))
        return real_text(pid, items, names)

    monkeypatch.setattr(punch_service, "_render_punchlist_text", text_spy)
    service = punch_service.PunchListService(MagicMock())
    service.repo = SimpleNamespace(all_for_project=AsyncMock(return_value=[item]))
    monkeypatch.setattr(service, "resolve_party_names", AsyncMock(return_value={}))

    pdf = await service.export_pdf(project_id)

    assert pdf.startswith(b"%PDF")
    assert seen, "no PDF renderer ran"
    assert threading.get_ident() not in seen
    (copies,) = handed or text_calls
    (copy,) = copies
    assert not isinstance(copy, PunchItem)
    for field in (
        "title",
        "description",
        "status",
        "priority",
        "category",
        "trade",
        "due_date",
        "resolution_notes",
        "photos",
        "metadata_",
        "page",
        "location_x",
        "location_y",
        "reopen_history",
    ):
        assert getattr(copy, field) == getattr(item, field), field
    # The text form both renderers share reads the copy exactly as it read the row.
    assert real_text(project_id, [copy], {}) == real_text(project_id, [item], {})
