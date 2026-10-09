# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Excel (.xlsx) and CSV BOQ importer.

Generic spreadsheet ingester with three classification heuristics on top
of the column-alias mapper:

* **NRM** - UK New Rules of Measurement. Detects element codes like
  ``2.6.1`` and section headers like ``Element 2 - Substructure``.
* **MasterFormat** - US CSI MasterFormat. Detects 6-digit codes like
  ``03 30 00`` and division headers like
  ``Division 03 - Cast-in-Place Concrete``.
* **Generic** - anything else goes into ``classification["code"]``.

Epic I3 (refactor) keeps the parser pure (no DB I/O, no FastAPI types);
all the per-row error reporting, dry-run handling and inline validation
the route used to do inline now live in the dispatcher route. Epics
I9 / I10 wire in the NRM and MasterFormat division detectors as
:func:`_infer_classification`.
"""

from __future__ import annotations

import csv
import io
import itertools
import logging
import math
import re
import unicodedata
from typing import Any, ClassVar, Literal

from app.core.file_signature import detect as detect_signature
from app.core.sheet_header import HEADER_SEARCH_ROWS, find_header_row
from app.modules.boq.importers._base import (
    ImportedBOQ,
    ImportedPosition,
    ImporterParseError,
)
from app.modules.boq.importers._encoding import (
    decode_text_bytes,
    parse_numeric_cell,
    safe_float,
)
from app.modules.boq.importers.hungary_workbook import parse_hungarian_workbook
from app.modules.boq.roundtrip import ID_COLUMN_ALIASES, normalise_id
from app.modules.boq.units import is_lump_sum_unit

logger = logging.getLogger(__name__)


# ── Column alias map, tagged by language ───────────────────────────────────
#
# Canonical column → accepted header strings (lowercased), grouped by the
# language whose market writes them. Editing this map is the supported
# extension point for new locale variants (Polish ``Ilosc``, Italian
# ``Quantità`` etc. land here).
#
# It is a table keyed by language rather than a flat set of strings because a
# caller has to be able to ask which languages a header row can be read in,
# and deriving that back out of a flat set means guessing which market owns
# ``prezzo``. ``_COLUMN_ALIASES`` below is the union, computed rather than
# maintained, so the matcher keeps behaving exactly as it did.
#
# Every accented header carries its unaccented twin: exports strip diacritics
# often enough that a table holding only ``descrição`` reads nothing out of a
# file whose header says ``DESCRICAO``.
#
# Strings no single language owns live under ``"en"``: the abbreviations an
# export writes whatever its locale (``pos``, ``nr``, ``qty``) and the
# classification standards (``nrm``, ``csi``, ``masterformat``). The German
# ones stay German - ``kg`` is a DIN 276 Kostengruppe, not a kilogramme.
_HEADERS_BY_LANGUAGE: dict[str, dict[str, tuple[str, ...]]] = {
    "en": {
        "ordinal": (
            "pos",
            "pos.",
            "position",
            "ordinal",
            "nr",
            "nr.",
            "no",
            "no.",
            "ord",
            "item",
            "item no",
            "ref",
            "#",
            # South Asian and Commonwealth bills: "Sl. No.", "S. No.", "Sr. No.".
            "sl. no.",
            "s. no.",
            "sr. no.",
        ),
        "description": (
            "description",
            "desc",
            "text",
            "item description",
            "description of item",
            "description of work",
        ),
        "unit": ("unit", "uom", "unit of measure"),
        "quantity": ("quantity", "qty", "qty."),
        "unit_rate": ("unit rate", "rate", "unitrate", "unit price", "unit cost", "price"),
        "total": ("total", "amount", "subtotal", "sum", "total price"),
        # Note: ``"code"`` lives here in the ``classification`` group, not in
        # ``ordinal``. Spreadsheets that name their classification column
        # "Code" (NRM / MasterFormat exports) need that header to map to
        # ``classification`` so the I9 / I10 heuristics can infer ``nrm`` /
        # ``masterformat``.
        "classification": (
            "classification",
            "nrm",
            "code",
            "csi",
            "masterformat",
            "element",
            "division",
            "category",
            "trade",
            "cost code",
            "cost group",
            "class",
        ),
        # A bill that prices material and labour apart carries two rates and
        # two totals per line and no single rate at all. See
        # :func:`_combine_split_columns` for how they become the one rate.
        "material_rate": ("material rate", "material unit rate", "material unit price"),
        "labour_rate": ("labour rate", "labor rate", "labour unit rate", "labor unit rate", "labour unit price"),
        "material_total": ("material total", "material amount"),
        "labour_total": ("labour total", "labor total", "labour amount", "labor amount"),
    },
    "de": {
        "ordinal": ("oz", "pos.-nr.", "lv-pos."),
        "description": ("beschreibung", "leistung", "bezeichnung", "kurztext"),
        "unit": ("einheit", "me"),
        "quantity": ("menge",),
        "unit_rate": ("einheitspreis", "ep", "preis"),
        "total": ("gesamt", "gesamtpreis", "gp"),
        "classification": ("din 276", "din276", "kg"),
    },
    "es": {
        "description": ("descripción", "descripcion", "designación", "designacion"),
        "unit": ("unidad", "uds", "ud"),
        "quantity": ("cantidad", "cant", "cant."),
        "unit_rate": ("precio",),
        "total": ("importe",),
    },
    "fr": {
        "description": ("désignation", "designation"),
        "unit": ("unité",),
        "quantity": ("quantité",),
        "unit_rate": ("prix",),
        "total": ("montant", "prix total", "montant ht", "total ht"),
    },
    "it": {
        "description": ("descrizione",),
        "unit": ("unità", "u"),
        "quantity": ("quantità", "quantita"),
        "unit_rate": ("prezzo",),
        "total": ("importo", "totale", "importo totale"),
    },
    "pl": {
        "description": ("opis",),
        "unit": ("jed",),
        "quantity": ("ilość", "ilosc"),
        # Polish carried no rate header at all until the table was split by
        # language, which made a Polish bill import with every rate at zero.
        "unit_rate": ("cena jednostkowa", "cena jedn.", "cena"),
    },
    "ru": {
        "description": ("наименование",),
        "unit": ("ед", "ед."),
        "quantity": ("количество", "кол-во"),
        "unit_rate": ("цена",),
        "total": ("стоимость",),
    },
    "pt": {
        "ordinal": ("nº", "n°", "n.º", "ordem"),
        "description": (
            "descrição",
            "descricao",
            "discriminação",
            "discriminacao",
            "especificação",
            "especificacao",
            "serviço",
            "servico",
        ),
        "unit": ("unidade", "und", "unid", "unid.", "un"),
        "quantity": ("quantidade", "qtd", "qtde", "quant", "quant."),
        "unit_rate": (
            "preço unitário",
            "preco unitario",
            "preço unit.",
            "preco unit.",
            "valor unitário",
            "valor unitario",
            "custo unitário",
            "custo unitario",
        ),
        "total": ("valor total", "preço total", "preco total"),
        # Brazilian estimators label the classification column after the
        # reference they priced from, so these must not fall to ``ordinal``.
        "classification": ("sinapi", "código sinapi", "codigo sinapi", "nbr", "nbr 12721"),
    },
    "nl": {
        "ordinal": ("post", "postnr", "postnr.", "volgnr", "volgnr."),
        "description": ("omschrijving", "beschrijving"),
        "unit": ("eenheid", "eenh", "eenh."),
        "quantity": ("hoeveelheid", "aantal", "hvh"),
        "unit_rate": ("eenheidsprijs", "prijs per eenheid", "prijs"),
        "total": ("totaal", "totaalbedrag", "bedrag"),
    },
    "cs": {
        "ordinal": ("poř.", "por.", "poř. č.", "por. c.", "p.č.", "p.c.", "pol."),
        "description": ("popis", "název", "nazev", "popis položky", "popis polozky"),
        "unit": ("mj", "m.j.", "měrná jednotka", "merna jednotka", "jednotka"),
        "quantity": ("množství", "mnozstvi", "výměra", "vymera"),
        "unit_rate": ("jednotková cena", "jednotkova cena", "j. cena", "cena"),
        "total": ("celkem", "cena celkem", "celková cena", "celkova cena"),
    },
    "sk": {
        "ordinal": ("por.", "por. č.", "p. č.", "p. c."),
        "description": ("popis", "názov", "nazov", "popis položky", "popis polozky"),
        "unit": ("mj", "m.j.", "merná jednotka", "merna jednotka", "jednotka"),
        "quantity": ("množstvo", "mnozstvo", "výmera", "vymera"),
        "unit_rate": ("jednotková cena", "jednotkova cena", "cena"),
        "total": ("spolu", "celkom", "cena spolu"),
    },
    "tr": {
        "ordinal": ("sıra", "sira", "sıra no", "sira no", "poz", "poz no"),
        "description": (
            "tanım",
            "tanim",
            "iş kalemi",
            "is kalemi",
            "açıklama",
            "aciklama",
            "imalatın cinsi",
            "imalatin cinsi",
        ),
        "unit": ("birim", "ölçü birimi", "olcu birimi"),
        "quantity": ("miktar", "metraj"),
        "unit_rate": ("birim fiyat", "birim fiyatı", "birim fiyati"),
        "total": ("tutar", "toplam", "toplam tutar"),
    },
    # Hungarian költségvetés. The item number ("Tételszám") is the norm or
    # catalogue code of the line, 21-003-5.1.1 and the like, not its running
    # number, so it is read as the classification and "Ssz." stays the
    # ordinal. Every priced line is quoted as material (anyag) plus fee (díj),
    # which a two-row header writes as "Egységár" over "Anyag | Díj".
    "hu": {
        "ordinal": ("sorszám", "sorszam", "ssz", "ssz.", "s.sz.", "sorsz."),
        "description": ("megnevezés", "megnevezes", "tétel szövege", "tetel szovege", "leírás", "leiras"),
        "unit": (
            "egység",
            "egyseg",
            "m.e.",
            "mennyiségi egység",
            "mennyisegi egyseg",
            "mértékegység",
            "mertekegyseg",
            "m.egys.",
        ),
        "quantity": ("mennyiség", "mennyiseg", "menny.", "menny"),
        "unit_rate": ("egységár", "egysegar", "egység ár", "egyseg ar"),
        "total": ("összesen", "osszesen", "összeg", "osszeg", "mindösszesen", "mindosszesen"),
        "classification": ("tételszám", "tetelszam", "tétel szám", "tetel szam", "normaszám", "normaszam"),
        "material_rate": ("anyag egységár", "anyag egysegar", "anyag egységára", "anyag egysegara"),
        "labour_rate": ("díj egységár", "dij egysegar", "munkadíj egységár", "munkadij egysegar"),
        "material_total": ("anyag összesen", "anyag osszesen", "nettó anyag összesen", "netto anyag osszesen"),
        "labour_total": (
            "díj összesen",
            "dij osszesen",
            "munkadíj összesen",
            "munkadij osszesen",
            "nettó díj összesen",
            "netto dij osszesen",
        ),
    },
    "ro": {
        "ordinal": ("nr. crt.", "nr crt", "crt.", "poz."),
        "description": ("denumire", "denumire lucrare", "denumirea lucrării", "denumirea lucrarii", "descriere"),
        "unit": ("um", "u.m.", "unitate", "unitate de măsură", "unitate de masura"),
        "quantity": ("cantitate",),
        "unit_rate": ("preț unitar", "pret unitar"),
        "total": ("valoare", "valoare totală", "valoare totala"),
    },
    "bg": {
        "ordinal": ("№", "поз", "поз."),
        "description": (
            "описание", "вид работа", "видове работи", "видове смр", "смр",
            "наименование",
            # Real КСС layouts: "Вид на СМР" (образец), "Вид дейност"
            # (ценово предложение) name the description column.
            "вид на смр", "вид дейност", "вид дейности", "вид на дейността",
            "вид на работите", "наименование на смр",
            # "Описание на строително-монтажни работи" and its short forms.
            "описание на смр", "описание на строително-монтажни работи",
            "описание на работите", "описание на дейността",
            "описание на дейностите", "наименование на работите",
            "наименование на дейностите",
        ),
        "unit": ("мярка", "ед. мярка", "един. мярка", "ед.м.", "единица мярка", "мерна единица"),
        "quantity": ("количество", "количество общо", "кол-во", "кол.", "к-во"),
        "unit_rate": (
            "ед. цена", "единична цена",
            # Plural and VAT-qualified forms used in tender КСС files:
            # "единични цени в EUR", "Ед. цена в евро без ДДС".
            "единични цени", "единични цени в", "единична цена в евро",
            "единична цена в лв", "единична цена без ддс",
            "ед. цена без ддс", "ед. цена в евро без ддс",
            "ед. цена в лв без ддс", "цена за единица", "цена на единица",
        ),
        "total": (
            "стойност", "обща стойност", "общо",
            # "Обща цена в евро без ДДС" and its variants.
            "обща цена", "обща цена в евро", "обща цена в лв",
            "обща цена без ддс", "обща цена в евро без ддс",
            "стойност без ддс", "стойност в евро", "обща стойност без ддс",
        ),
    },
    "el": {
        "ordinal": ("α/α", "αα"),
        "description": ("περιγραφή", "περιγραφη", "είδος εργασίας", "ειδος εργασιας", "ονομασία", "ονομασια"),
        "unit": ("μονάδα", "μοναδα", "μονάδα μέτρησης", "μοναδα μετρησης", "μ.μ."),
        "quantity": ("ποσότητα", "ποσοτητα"),
        "unit_rate": ("τιμή μονάδας", "τιμη μοναδας", "τιμή", "τιμη"),
        "total": ("σύνολο", "συνολο", "δαπάνη", "δαπανη"),
    },
    "sv": {
        "ordinal": ("post", "postnr"),
        "description": ("beskrivning", "benämning", "benamning"),
        "unit": ("enhet", "enh", "enh."),
        "quantity": ("mängd", "mangd", "antal"),
        "unit_rate": ("à-pris", "a-pris", "enhetspris"),
        "total": ("summa", "belopp", "totalt"),
    },
    "no": {
        "ordinal": ("post", "postnr", "postnr."),
        "description": ("beskrivelse", "betegnelse"),
        "unit": ("enhet", "enh", "enh."),
        "quantity": ("mengde", "antall"),
        "unit_rate": ("enhetspris", "pris"),
        "total": ("sum", "beløp", "belop", "totalt"),
    },
    "da": {
        "ordinal": ("post", "postnr", "løbenr", "lobenr"),
        "description": ("beskrivelse", "betegnelse", "ydelse"),
        "unit": ("enhed", "enh", "enh."),
        "quantity": ("mængde", "maengde", "antal"),
        "unit_rate": ("enhedspris", "pris"),
        "total": ("sum", "beløb", "belob", "i alt"),
    },
    "fi": {
        "ordinal": ("nro", "nro.", "n:o"),
        "description": ("kuvaus", "selite", "nimike", "työn kuvaus", "tyon kuvaus"),
        "unit": ("yksikkö", "yksikko", "yks", "yks."),
        "quantity": ("määrä", "maara"),
        "unit_rate": ("yksikköhinta", "yksikkohinta", "yks.hinta", "hinta"),
        "total": ("yhteensä", "yhteensa", "summa", "kokonaishinta"),
    },
    "uk": {
        "ordinal": ("№ з/п", "поз", "поз."),
        "description": ("найменування", "опис", "найменування робіт"),
        "unit": ("од", "од.", "од. вим.", "одиниця виміру", "одиниця"),
        "quantity": ("кількість", "к-ть"),
        "unit_rate": ("ціна", "ціна за одиницю", "вартість одиниці"),
        "total": ("сума", "вартість", "загальна вартість"),
    },
    "ja": {
        "ordinal": ("番号", "項番"),
        "description": ("名称", "工種", "摘要", "工事内容", "説明", "内容"),
        "unit": ("単位",),
        "quantity": ("数量",),
        "unit_rate": ("単価",),
        "total": ("金額", "合計"),
    },
    "ko": {
        "ordinal": ("번호", "순번", "연번"),
        "description": ("품명", "공종", "내역", "설명", "공사명"),
        "unit": ("단위",),
        "quantity": ("수량",),
        "unit_rate": ("단가",),
        "total": ("금액", "합계"),
    },
    "zh": {
        "ordinal": ("序号", "编号"),
        "description": ("名称", "项目名称", "工作内容", "描述", "项目描述"),
        "unit": ("单位", "计量单位"),
        "quantity": ("数量", "工程量"),
        "unit_rate": ("单价", "综合单价"),
        "total": ("合价", "金额", "合计"),
    },
    "ar": {
        "ordinal": ("رقم", "الرقم", "التسلسل", "رقم البند"),
        "description": ("الوصف", "وصف", "البيان", "وصف الأعمال", "البند"),
        "unit": ("الوحدة", "وحدة", "وحدة القياس"),
        "quantity": ("الكمية", "كمية"),
        "unit_rate": ("سعر الوحدة", "السعر", "سعر"),
        "total": ("الإجمالي", "الاجمالي", "المجموع"),
    },
    "he": {
        "ordinal": ("מס'", "מספר", "סעיף"),
        "description": ("תיאור", "תאור", "פירוט"),
        "unit": ("יחידה", "יח'", "יחידת מידה"),
        "quantity": ("כמות",),
        "unit_rate": ("מחיר יחידה", "מחיר"),
        "total": ('סה"כ', "סהכ", "סך הכל", "סכום"),
    },
    "id": {
        "ordinal": ("nomor", "urut", "no. urut"),
        "description": ("uraian", "uraian pekerjaan", "deskripsi", "jenis pekerjaan"),
        "unit": ("satuan", "sat", "sat."),
        # ``jumlah`` is deliberately absent. Indonesian bills head both the
        # quantity column and the money column with it, so accepting it makes
        # one of the two read as the other; ``volume`` and ``jumlah harga``
        # are the spellings that say which is meant.
        "quantity": ("volume", "vol.", "kuantitas", "banyaknya"),
        "unit_rate": ("harga satuan", "harga"),
        "total": ("jumlah harga", "total harga", "jumlah biaya"),
    },
    "vi": {
        "ordinal": ("stt", "số tt", "so tt"),
        "description": (
            "nội dung công việc",
            "noi dung cong viec",
            "tên công việc",
            "ten cong viec",
            "mô tả",
            "mo ta",
            "diễn giải",
            "dien giai",
        ),
        "unit": ("đơn vị", "don vi", "đơn vị tính", "don vi tinh", "đvt", "dvt"),
        "quantity": ("khối lượng", "khoi luong", "số lượng", "so luong"),
        "unit_rate": ("đơn giá", "don gia"),
        "total": ("thành tiền", "thanh tien", "tổng cộng", "tong cong"),
    },
    # Croatian troškovnik. "Jed. mj." is the unit (jedinica mjere), and the
    # money columns usually carry the currency in brackets, "Jed. cijena
    # (EUR)", which the matcher strips before it looks the header up.
    "hr": {
        "ordinal": ("r.br.", "r. br.", "rbr", "rb", "red. br.", "redni broj", "br.", "br. stavke"),
        "description": ("opis", "opis stavke", "opis radova", "opis rada", "naziv", "naziv stavke"),
        "unit": ("jed. mj.", "jed.mj.", "j. mj.", "jm", "jedinica mjere", "mjerna jedinica"),
        "quantity": ("količina", "kolicina", "kol."),
        "unit_rate": ("jed. cijena", "jedinična cijena", "jedinicna cijena", "cijena"),
        "total": ("ukupno", "iznos", "ukupna cijena", "ukupni iznos"),
    },
    # Serbian writes both scripts, so each Latin spelling has its Cyrillic twin.
    "sr": {
        "ordinal": ("r.br.", "rb", "redni broj", "р.бр.", "рб", "редни број"),
        "description": ("opis", "opis pozicije", "naziv", "опис", "опис позиције", "назив"),
        "unit": ("jed. mere", "jedinica mere", "j.m.", "јед. мере", "јединица мере", "ј.м."),
        "quantity": ("količina", "kolicina", "количина"),
        "unit_rate": ("jed. cena", "jedinična cena", "jedinicna cena", "cena", "јед. цена", "јединична цена"),
        "total": ("ukupno", "iznos", "укупно", "износ"),
    },
    "sl": {
        "ordinal": ("zap. št.", "zap. st.", "zap.št.", "poz."),
        "description": ("opis", "opis postavke", "naziv"),
        "unit": ("enota", "enota mere", "em", "e.m."),
        "quantity": ("količina", "kolicina"),
        "unit_rate": ("cena na enoto", "cena/enoto", "enotna cena", "cena enote", "cena"),
        "total": ("skupaj", "vrednost", "znesek"),
    },
    "et": {
        "ordinal": ("jrk", "jrk nr", "jrk. nr"),
        "description": ("kirjeldus", "nimetus", "töö kirjeldus"),
        "unit": ("ühik", "uhik", "mõõtühik", "mootuhik"),
        "quantity": ("kogus", "maht"),
        "unit_rate": ("ühikuhind", "uhikuhind", "ühiku hind", "hind"),
        "total": ("kokku", "maksumus"),
    },
    "th": {
        "ordinal": ("ลำดับ", "ลำดับที่"),
        "description": ("รายการ", "รายละเอียด"),
        "unit": ("หน่วย",),
        "quantity": ("จำนวน", "ปริมาณ"),
        "unit_rate": ("ราคาต่อหน่วย", "ราคา/หน่วย"),
        "total": ("รวมเงิน", "จำนวนเงิน"),
    },
    "hi": {
        "ordinal": ("क्रम सं.", "क्रमांक", "क्र.सं."),
        "description": ("विवरण", "कार्य का विवरण"),
        "unit": ("इकाई",),
        "quantity": ("मात्रा",),
        "unit_rate": ("दर",),
        "total": ("राशि", "कुल राशि"),
    },
    "bn": {
        "ordinal": ("ক্রমিক নং", "ক্রম"),
        "description": ("বিবরণ", "কাজের বিবরণ"),
        "unit": ("একক",),
        "quantity": ("পরিমাণ",),
        "unit_rate": ("দর", "একক দর"),
        "total": ("মোট", "মোট টাকা"),
    },
    # Urdu "شرح" means rate but Persian "شرح" means description, so Urdu
    # carries only the loanword for rate and the shared spelling is Persian's.
    "ur": {
        "ordinal": ("نمبر شمار",),
        "description": ("تفصیل",),
        "unit": ("اکائی", "یونٹ"),
        "quantity": ("مقدار",),
        "unit_rate": ("ریٹ", "فی یونٹ ریٹ"),
        # Bare "رقم" is Arabic for the item number, so Urdu keeps only the
        # spelling that says "total amount".
        "total": ("کل رقم",),
    },
    "fa": {
        "ordinal": ("ردیف",),
        "description": ("شرح", "شرح عملیات", "شرح کار"),
        "unit": ("واحد",),
        # Persian and Urdu share the spelling and the meaning here.
        "quantity": ("مقدار",),
        "unit_rate": ("بهای واحد", "فی"),
        "total": ("بهای کل", "مبلغ"),
    },
    "fil": {
        "ordinal": ("blg.",),
        "description": ("paglalarawan",),
        "unit": ("yunit",),
        "quantity": ("dami",),
        "unit_rate": ("presyo bawat yunit", "halaga bawat yunit"),
        "total": ("kabuuan", "kabuuang halaga"),
    },
    "kk": {
        "ordinal": ("р/с", "№ р/с"),
        "description": ("атауы", "жұмыстардың атауы"),
        "unit": ("өлшем бірлігі", "өлш. бір."),
        "quantity": ("саны", "көлемі"),
        "unit_rate": ("бірлік бағасы", "бағасы"),
        "total": ("сомасы", "құны"),
    },
    "ky": {
        "description": ("аталышы", "иштердин аталышы"),
        "unit": ("өлчөө бирдиги",),
        # Kazakh and Kyrgyz share the spelling and the meaning here.
        "quantity": ("саны",),
        "unit_rate": ("баасы", "бирдик баасы"),
        "total": ("суммасы",),
    },
    "uz": {
        "ordinal": ("t/r",),
        "description": ("nomi", "ishlar nomi"),
        "unit": ("o'lchov birligi", "o‘lchov birligi", "olchov birligi"),
        "quantity": ("miqdori", "soni"),
        "unit_rate": ("narxi", "birlik narxi"),
        "total": ("qiymati", "jami"),
    },
    "mn": {
        "ordinal": ("д/д",),
        "description": ("ажлын нэр", "нэр"),
        "unit": ("хэмжих нэгж", "нэгж"),
        "quantity": ("тоо хэмжээ",),
        "unit_rate": ("нэгж үнэ",),
        "total": ("нийт үнэ", "дүн"),
    },
}


# The four columns a bill row cannot be read without. A language that names
# fewer than these cannot carry a bill on its own, however many other headers
# it declares, so it has no business being listed as supported.
_MANDATORY_COLUMNS: tuple[str, ...] = ("description", "unit", "quantity", "unit_rate")

# Canonical column order for the flattened map. ``position_id`` is prepended
# by the builder and comes first: an exported "Position ID" header maps there,
# never to ``ordinal``. A blank cell -> new row; a value belonging to the
# target BOQ -> update in place (GitHub #360).
_CANONICAL_COLUMNS: tuple[str, ...] = (
    "ordinal",
    "description",
    "unit",
    "quantity",
    "unit_rate",
    "total",
    "classification",
)


def _languages_missing_mandatory_columns(
    table: dict[str, dict[str, tuple[str, ...]]],
) -> dict[str, tuple[str, ...]]:
    """Report which mandatory columns each language fails to name.

    Args:
        table: A language-tagged header table shaped like
            :data:`_HEADERS_BY_LANGUAGE`.

    Returns:
        Language code -> the mandatory columns it leaves empty or omits.
        Empty when every language is complete.
    """
    holes: dict[str, tuple[str, ...]] = {}
    for language, headers in table.items():
        missing = tuple(column for column in _MANDATORY_COLUMNS if not headers.get(column))
        if missing:
            holes[language] = missing
    return holes


def _build_column_aliases(
    table: dict[str, dict[str, tuple[str, ...]]],
) -> dict[str, frozenset[str]]:
    """Flatten the language-tagged table into the canonical alias map.

    Args:
        table: A language-tagged header table shaped like
            :data:`_HEADERS_BY_LANGUAGE`.

    Returns:
        Canonical column -> every accepted header string for it, across all
        languages, with ``position_id`` seeded from
        :data:`~app.modules.boq.roundtrip.ID_COLUMN_ALIASES`.
    """
    merged: dict[str, set[str]] = {}
    for headers in table.values():
        for canonical, words in headers.items():
            merged.setdefault(canonical, set()).update(words)
    aliases: dict[str, frozenset[str]] = {"position_id": ID_COLUMN_ALIASES}
    for canonical in _CANONICAL_COLUMNS:
        aliases[canonical] = frozenset(merged.pop(canonical, set()))
    for canonical in sorted(merged):
        aliases[canonical] = frozenset(merged[canonical])
    return aliases


_COLUMN_ALIASES: dict[str, frozenset[str]] = _build_column_aliases(_HEADERS_BY_LANGUAGE)

SUPPORTED_HEADER_LANGUAGES: frozenset[str] = frozenset(_HEADERS_BY_LANGUAGE)
"""Languages whose spreadsheet header row this importer reads natively.

Computed from :data:`_HEADERS_BY_LANGUAGE`, never written out by hand, so a
caller deciding whether a national profile can claim native spreadsheet
import reads what the table actually holds rather than what a second list
once said it held. Membership means the language names at least
:data:`_MANDATORY_COLUMNS`.

Note that this covers the header row only. A market whose bills are not
tables with a header row at all (the Hungarian workbooks, whose item code is
composed down a heading tree across nine columns) needs a profile of its own
regardless of what this set says.
"""


# ── Label normalisation ─────────────────────────────────────────────────────
#
# A header is written a dozen ways for the same column: "Jed. mj.", "JED MJ",
# "Jed.mj.", "Jed. cijena (EUR)", "KOLIČINA", "Kolicina". The table above
# holds one or two spellings per column, so the matcher reduces both sides to
# a key that ignores case, diacritics, punctuation, spacing and a bracketed or
# trailing currency before it compares them.

# Letters NFKD leaves alone because Unicode does not treat them as a base
# letter plus an accent. Without these, "Količina" and "Kolicina" meet but
# "Đ" (Croatian), "Ł" (Polish), "Ø"/"Æ" (Danish, Norwegian) and "ı"
# (Turkish) never reach their unaccented twins.
_TRANSLITERATE: dict[int, str] = str.maketrans(
    {"đ": "d", "ł": "l", "ø": "o", "æ": "ae", "œ": "oe", "ß": "ss", "ı": "i", "þ": "th", "ð": "d"}
)

# Combining marks are dropped only above this code point's scripts (Latin,
# Greek, Cyrillic, Armenian). Thai tone marks, Devanagari and Bengali vowel
# signs and Arabic hamza are combining marks too, but there they change the
# word, so dropping them would make two different headers one.
_STRIP_MARKS_BELOW = 0x0590

_BRACKETED = re.compile(r"\([^)]*\)|\[[^\]]*\]|\{[^}]*\}")

# A Bulgarian header typed on the wrong keyboard layout slips a Latin twin
# into a Cyrillic word ("Eд. Мярка" starts with a Latin E, U+0045). The fold
# is applied per word and only when the word already carries real Cyrillic,
# so a pure-Latin label ("Netto", "Tax") is never rewritten.
_HOMOGLYPH_TO_CYRILLIC: dict[int, str] = str.maketrans(
    {
        "a": "а", "b": "в", "c": "с", "e": "е", "h": "н", "k": "к",
        "m": "м", "o": "о", "p": "р", "t": "т", "x": "х", "y": "у",
    }
)
_CYRILLIC_WORD = re.compile(r"[а-яёіїєґ]")


def _currency_codes() -> frozenset[str]:
    """Lowercased ISO 4217 codes the platform knows, for stripping from headers."""
    from app.core.currency_registry import CURRENCIES

    return frozenset(code.lower() for code in CURRENCIES)


_CURRENCY_CODES: frozenset[str] = _currency_codes()


def normalise_label(text: str) -> str:
    """Reduce a header or row label to its comparison form.

    Lowercases, drops bracketed parts ("Jed. cijena (EUR)"), strips accents
    from Latin, Greek and Cyrillic letters, turns every run of punctuation
    and spacing into one space and drops a trailing currency code or sign
    ("Iznos EUR", "Ukupno €"). The words stay separated by single spaces.

    Args:
        text: The raw cell text.

    Returns:
        The normalised label, possibly empty.
    """
    lowered = _BRACKETED.sub(" ", str(text).lower()).translate(_TRANSLITERATE)
    kept: list[str] = []
    base = 0
    for char in unicodedata.normalize("NFKD", lowered):
        if unicodedata.combining(char):
            if base < _STRIP_MARKS_BELOW:
                continue
            kept.append(char)
            continue
        base = ord(char)
        kept.append(char if char.isalnum() else " ")
    words = unicodedata.normalize("NFC", "".join(kept)).split()
    words = [
        word.translate(_HOMOGLYPH_TO_CYRILLIC) if _CYRILLIC_WORD.search(word) else word
        for word in words
    ]
    if len(words) > 1 and words[-1] in _CURRENCY_CODES:
        words.pop()
    return " ".join(words)


def _label_key(text: str) -> str:
    """The normalised label with the spaces removed too ("r br" == "rbr")."""
    return normalise_label(text).replace(" ", "")


def _build_normalised_index(aliases: dict[str, frozenset[str]]) -> dict[str, str]:
    """Key every alias by :func:`_label_key`, the first canonical column winning.

    The collision test in ``tests/unit/test_boq_spreadsheet_header_languages``
    holds that no two canonical columns share a key, so "first wins" never
    decides anything in practice; it only keeps ``position_id`` ahead of
    ``ordinal`` if that test is ever broken.
    """
    index: dict[str, str] = {}
    for canonical, words in aliases.items():
        for word in words:
            key = _label_key(word)
            if key:
                index.setdefault(key, canonical)
    return index


_NORMALISED_COLUMN_INDEX: dict[str, str] = _build_normalised_index(_COLUMN_ALIASES)


def _match_column(header: str) -> str | None:
    """Match a header string to a canonical column name using the alias map.

    The exact lowercased spelling is tried first, so a symbol-only header such
    as ``#`` (which normalises to nothing) still matches, then the normalised
    key, which is what reads "Jed. mj." and "KOLICINA (m3)".
    """
    lowered = header.strip().lower()
    for canonical, aliases in _COLUMN_ALIASES.items():
        if lowered in aliases:
            return canonical
    key = _label_key(header)
    return _NORMALISED_COLUMN_INDEX.get(key) if key else None


# ── Split rates, header language, two-row headers ──────────────────────────

# The split columns and the single column each one feeds. A Hungarian bill
# prices every line as material plus fee and has no single rate column at
# all, so reading only ``unit_rate`` imported every line at zero, and reading
# one half as the rate would halve the bill.
_SPLIT_COLUMNS: dict[str, str] = {
    "material_rate": "unit_rate",
    "labour_rate": "unit_rate",
    "material_total": "total",
    "labour_total": "total",
}

# Languages whose CSV exports come out of Excel in Windows-1250. That code
# page decodes as Windows-1252 without an error, so nothing fails: the
# Hungarian ő and ű quietly become õ and û. The header row says which market
# wrote the file, and its words are plain enough to read in either code page.
_CP1250_LANGUAGES: frozenset[str] = frozenset({"hu", "cs", "sk", "pl", "hr", "sl", "sr", "ro"})

_LANGUAGE_HEADER_KEYS: dict[str, frozenset[str]] = {
    language: frozenset(_label_key(word) for words in headers.values() for word in words)
    for language, headers in _HEADERS_BY_LANGUAGE.items()
}


def header_language(header: tuple[Any, ...] | list[Any] | None) -> str | None:
    """The language whose table names the most cells of a header row.

    Args:
        header: The header row's cells.

    Returns:
        The language code, or ``None`` when no cell is a known header. A tie
        goes to the language listed first in :data:`_HEADERS_BY_LANGUAGE`.
    """
    keys = [_label_key(str(cell)) for cell in header or () if cell is not None and str(cell).strip()]
    best: str | None = None
    best_count = 0
    for language, known in _LANGUAGE_HEADER_KEYS.items():
        count = sum(1 for key in keys if key and key in known)
        if count > best_count:
            best, best_count = language, count
    return best


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _compose_two_row_header(header: tuple[Any, ...], below: tuple[Any, ...]) -> tuple[str, ...] | None:
    """Read a header written over two rows, or ``None`` when it is one row.

    Split bills head their money columns twice: "Egységár" merged across the
    two columns under it, which say "Anyag" and "Díj". A merged cell holds its
    text in the first column only, so the parent label is carried right across
    the empty cells of its span, and each column under it is named "<sub>
    <parent>", which is what the alias table spells ("anyag egységár").

    The composition is taken only when the row below holds no number and the
    composed header names more columns than the first row alone. A first data
    row that happens to be a text-only section heading composes into nothing
    the table knows and is left alone.
    """
    if not any(_cell_text(cell) for cell in below):
        return None
    if any(not math.isnan(safe_float(cell, default=math.nan)) for cell in below):
        return None
    width = max(len(header), len(below))
    composed: list[str] = []
    parent = ""
    for index in range(width):
        top = _cell_text(header[index]) if index < len(header) else ""
        sub = _cell_text(below[index]) if index < len(below) else ""
        if top:
            parent = top if sub else ""
        elif not sub:
            parent = ""
        if sub and parent:
            composed.append(f"{sub} {parent}")
        else:
            composed.append(top or sub)
    known_before = sum(1 for cell in header if _cell_text(cell) and _match_column(_cell_text(cell)))
    known_after = sum(1 for cell in composed if cell and _match_column(cell))
    return tuple(composed) if known_after > known_before else None


def _map_columns(header: tuple[Any, ...]) -> dict[int, str]:
    """Column index -> canonical column for every header cell the table knows."""
    column_map: dict[int, str] = {}
    for index, cell in enumerate(header):
        text = _cell_text(cell)
        canonical = _match_column(text) if text else None
        if canonical:
            column_map[index] = canonical
    return column_map


def _combine_split_columns(row: dict[str, Any]) -> dict[str, Any]:
    """Fold material and labour columns into ``unit_rate`` and ``total``.

    Only when the file has no single column of its own for the target: a
    bill that carries "Egységár" beside the split keeps its own figure. A
    blank half counts as zero (a fee-only line is a correct line), an unread
    half is handed on as text so the row reports it rather than importing at
    a guessed rate, and a row where both halves are blank stays unpriced, so
    a heading or a total line is still recognised as one.
    """
    for target in ("unit_rate", "total"):
        halves = [key for key, feeds in _SPLIT_COLUMNS.items() if feeds == target and key in row]
        if not halves or not _is_blank_value(row.get(target)):
            continue
        amount = 0.0
        filled = False
        for key in halves:
            value = row[key]
            if _is_blank_value(value):
                continue
            parsed, error = parse_numeric_cell(value)
            if error is not None or parsed is None:
                row[target] = value
                break
            amount += parsed
            filled = True
        else:
            if filled:
                row[target] = amount
    return row


def _is_blank_value(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _report_mapping(column_map: dict[int, str]) -> dict[str, str]:
    """The mapping as the import dialog shows it: a split column under the column it feeds."""
    return {str(index): _SPLIT_COLUMNS.get(canonical, canonical) for index, canonical in column_map.items()}


def _detect_file_format(content_head: bytes) -> Literal["xlsx", "csv", "parquet", "unknown"]:
    """Identify an upload by its magic bytes (BUG-UPLOAD01 from the legacy code).

    A ``.exe`` renamed to ``.xlsx`` would otherwise be handed to
    ``openpyxl`` - best case a parse exception, worst case the bytes
    land in our buffers + logs before we error.
    """
    if not content_head:
        return "unknown"
    sig = detect_signature(content_head)
    if sig == "zip":  # XLSX = OOXML zip
        return "xlsx"
    if content_head[:4] == b"PAR1":
        return "parquet"
    if b"\x00" in content_head:
        return "unknown"
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            decoded = content_head.decode(encoding)
        except UnicodeDecodeError:
            continue
        if any(sep in decoded for sep in (",", ";", "\t", "|", "\n")):
            return "csv"
    return "unknown"


# ── Classification heuristics (Epics I9 + I10) ─────────────────────────────


# NRM codes: ``N.N.N`` or ``N.N`` (e.g. ``2.6.1``, ``2.6``). NRM 1 / NRM 2
# tops out at four levels but two/three are the common case in tender
# documents.
_NRM_CODE_RE = re.compile(r"^(\d{1,2}\.){1,3}\d{1,2}$")

# NRM element header text e.g. ``"Element 2 - Substructure"``,
# ``"Group element 2.6 - External walls"``.
_NRM_HEADER_RE = re.compile(r"^(group\s+)?element\s+(\d{1,2}(?:\.\d{1,2})*)\b", re.IGNORECASE)

# MasterFormat: ``XX XX XX`` or ``XX.XX.XX`` or ``XX-XX-XX`` (2-2-2 digits).
# Sub-codes ``XX XX XX.XX`` are allowed.
_MASTERFORMAT_CODE_RE = re.compile(r"^(\d{2})[\s.\-](\d{2})[\s.\-](\d{2})(?:\.(\d{2}))?$")

# MasterFormat division header text e.g. ``"Division 03 - Concrete work"``,
# ``"03 30 00 Cast-in-Place Concrete"``.
_MASTERFORMAT_HEADER_RE = re.compile(r"^division\s+(\d{2})\b", re.IGNORECASE)


def _infer_classification(
    code_text: str,
    description: str,
) -> dict[str, Any]:
    """Heuristic classification from a raw code cell + description.

    Tries NRM and MasterFormat patterns; anything else falls through to
    ``{"code": code_text}`` (the historic generic behaviour).
    """
    code = code_text.strip()
    desc = (description or "").strip()
    classification: dict[str, Any] = {}

    # NRM element header in the description ("Element 2 - Substructure").
    m = _NRM_HEADER_RE.match(desc)
    if m:
        classification["nrm"] = m.group(2)
    # NRM code pattern in the code cell ("2.6.1").
    if code and _NRM_CODE_RE.match(code):
        classification["nrm"] = code

    # MasterFormat 6-digit code in the code cell ("03 30 00").
    m = _MASTERFORMAT_CODE_RE.match(code) if code else None
    if m:
        # Normalise to spaced form "XX XX XX[.XX]".
        parts = [m.group(1), m.group(2), m.group(3)]
        mf = " ".join(parts)
        if m.group(4):
            mf = f"{mf}.{m.group(4)}"
        classification["masterformat"] = mf

    # MasterFormat division header in the description ("Division 03 -").
    m = _MASTERFORMAT_HEADER_RE.match(desc)
    if m:
        # Pad to canonical 6-digit form for downstream rules.
        div = m.group(1)
        # If the description contains a fuller code further along, keep it,
        # else stub the level-2 + level-3 to ``00``.
        if "masterformat" not in classification:
            classification["masterformat"] = f"{div} 00 00"

    # Fallback: stash the raw code so the editor can show it. Skip if we
    # already mapped it to a structured field above.
    if code and "nrm" not in classification and "masterformat" not in classification:
        classification["code"] = code

    return classification


# ── Row parsing helpers ─────────────────────────────────────────────────────


_CSV_DELIMITERS: tuple[str, ...] = (";", "\t", ",", "|")


def _sniff_delimiter(text: str) -> str:
    """The delimiter that splits the most lines into the same number of fields.

    ``csv.Sniffer`` reads the whole sample, title lines included, and a
    semicolon file whose numbers carry decimal commas gives it two plausible
    answers. Counting fields line by line with the csv reader itself (so a
    quoted "125,5" stays one field) and taking the delimiter whose most common
    field count is shared by the most lines picks the one the table is laid
    out in, whatever sits above it. A tie goes to the earlier delimiter in
    :data:`_CSV_DELIMITERS`, semicolon first, which is what Excel writes in
    every locale that uses the decimal comma.
    """
    lines = [line for line in text[:16384].splitlines()[:60] if line.strip()]
    best, best_score = ",", 0
    for delimiter in _CSV_DELIMITERS:
        counts: dict[int, int] = {}
        for fields in csv.reader(lines, delimiter=delimiter):
            if len(fields) > 1:
                counts[len(fields)] = counts.get(len(fields), 0) + 1
        score = max(counts.values(), default=0)
        if score > best_score:
            best, best_score = delimiter, score
    return best


def _locate_header(
    rows_iter: Any,
) -> tuple[tuple[Any, ...] | None, int, Any]:
    """Find the header row and read a second header row under it if there is one.

    Returns ``(header, number of the last header row, rows under it)``.
    """
    raw_headers, header_number, rows = find_header_row(iter(rows_iter), _match_column)
    if not raw_headers:
        return raw_headers, header_number, rows
    below = next(rows, None)
    if below is None:
        return tuple(raw_headers), header_number, rows
    composed = _compose_two_row_header(tuple(raw_headers), tuple(below))
    if composed is not None:
        return composed, header_number + 1, rows
    return tuple(raw_headers), header_number, itertools.chain([below], rows)


def _decode_csv(content_bytes: bytes) -> tuple[str, str]:
    """Decode a CSV upload, reading Windows-1250 files as Windows-1250.

    :func:`decode_text_bytes` answers Windows-1252 for any single-byte file,
    and that answer is wrong for a Central European export without being an
    error. When the first answer is a single-byte code page and the header
    row reads in a language of :data:`_CP1250_LANGUAGES`, the bytes are
    decoded again as Windows-1250. Kept to this reader: BC3 and the other
    text formats share :data:`DEFAULT_ENCODINGS` and are Western by
    convention.
    """
    text, encoding = decode_text_bytes(content_bytes)
    if encoding not in ("cp1252", "latin-1"):
        return text, encoding
    try:
        central = content_bytes.decode("cp1250")
    except UnicodeDecodeError:
        return text, encoding
    if central == text:
        return text, encoding
    delimiter = _sniff_delimiter(text)
    header, _, _ = find_header_row(csv.reader(io.StringIO(text), delimiter=delimiter), _match_column)
    if header_language(header) in _CP1250_LANGUAGES:
        return central, "cp1250"
    return text, encoding


def _parse_csv(content_bytes: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Decode + parse a CSV into canonical-key dicts and import metadata."""
    text, encoding = _decode_csv(content_bytes)
    delimiter = _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_headers, header_number, rows_iter = _locate_header(reader)
    if not raw_headers:
        raise ImporterParseError("CSV file is empty or has no header row")

    column_map = _map_columns(raw_headers)

    rows: list[dict[str, Any]] = []
    row_numbers: list[int] = []
    for line_number, raw_row in enumerate(rows_iter, start=header_number + 1):
        row: dict[str, Any] = {}
        for idx, val in enumerate(raw_row):
            canonical = column_map.get(idx)
            if canonical:
                row[canonical] = val.strip() if isinstance(val, str) else val
        if row:
            rows.append(_combine_split_columns(row))
            row_numbers.append(line_number)

    import_metadata = {
        "original_columns": [_cell_text(h) for h in raw_headers],
        "column_mapping": _report_mapping(column_map),
        "header_language": header_language(raw_headers),
        "encoding": encoding,
        "delimiter": delimiter,
        "total_rows": len(rows),
        "row_numbers": row_numbers,
    }
    return rows, import_metadata


def _parse_rows_from_csv(content_bytes: bytes) -> list[dict[str, Any]]:
    """Decode + parse a CSV into a list of canonical-key dicts."""
    rows, _ = _parse_csv(content_bytes)
    return rows


def _pick_item_sheet(workbook: Any) -> Any:
    """The worksheet the bill's lines are on.

    The active sheet when it carries a header naming a description and a
    quantity or a price. Otherwise the first sheet that does: exported bills
    open on a cover or a summary sheet ("Záradék", "Összesítő") and keep the
    lines on a later one, and reading only the active sheet found no rows at
    all. When no sheet qualifies the active sheet is returned as before, so
    the error the user sees is unchanged.
    """
    active = workbook.active
    candidates = [active] + [workbook[name] for name in workbook.sheetnames if workbook[name] is not active]
    for worksheet in candidates:
        if worksheet is None:
            continue
        header, _, _ = _locate_header(worksheet.iter_rows(max_row=HEADER_SEARCH_ROWS + 1, values_only=True))
        mapped = set(_map_columns(header or ()).values())
        priced = mapped & {"quantity", "unit_rate", *_SPLIT_COLUMNS}
        if "description" in mapped and priced:
            return worksheet
    return active


def _parse_rows_from_excel(
    content_bytes: bytes,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Read the item sheet of an .xlsx file into canonical-key dicts.

    The item sheet is the active one unless it carries no bill header, see
    :func:`_pick_item_sheet`. Returns ``(rows, import_metadata)``; metadata
    preserves the raw column ordering so a later export can round-trip back
    to the user's original spreadsheet layout, and ``row_numbers`` holds the
    sheet row each returned row came from, so a message about it names the
    row the user sees under a letterhead and past blank lines.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True)
    if wb.active is None:
        raise ImporterParseError("Excel file has no worksheets")
    ws = _pick_item_sheet(wb)

    sheet_names = wb.sheetnames

    raw_headers, header_number, rows_iter = _locate_header(ws.iter_rows(values_only=True))
    if not raw_headers:
        raise ImporterParseError("Excel file is empty or has no header row")

    original_columns = [str(h) if h is not None else "" for h in raw_headers]
    column_map = _map_columns(raw_headers)

    rows: list[dict[str, Any]] = []
    row_numbers: list[int] = []
    for sheet_row, raw_row in enumerate(rows_iter, start=header_number + 1):
        row: dict[str, Any] = {}
        for idx, val in enumerate(raw_row):
            canonical = column_map.get(idx)
            if canonical and val is not None:
                row[canonical] = val
        if row:
            rows.append(_combine_split_columns(row))
            row_numbers.append(sheet_row)
    item_sheet = ws.title
    wb.close()

    import_metadata = {
        "original_columns": original_columns,
        "column_mapping": _report_mapping(column_map),
        "header_language": header_language(raw_headers),
        "sheet_names": sheet_names,
        "item_sheet": item_sheet,
        "total_rows": len(rows),
        "row_numbers": row_numbers,
    }
    return rows, import_metadata


_TOTAL_ROW_DESCRIPTIONS = {
    "grand total",
    "total",
    "summe",
    "gesamt",
    "gesamtsumme",
    "subtotal",
    "zwischensumme",
    # Export artifacts of our own workbook - never re-imported as positions
    # on a round-trip (GitHub #360).
    "direct cost",
    "cost summary",
    "net total",
    "gross total",
}


# ── Total, tax and recap rows, tagged by language ──────────────────────────
#
# A national bill closes every section with a total line ("UKUPNO I. ...",
# "Summe Titel 01"), ends with a tax line and a grand total ("PDV 25 %",
# "SVEUKUPNO") and often repeats the section totals on a recap page
# ("REKAPITULACIJA"). None of them is work, and a total line has no unit,
# quantity or rate, so without this table each one imported as an empty
# section, and the recap page repeated the section ordinals.
#
# The phrases are matched on :func:`normalise_label`, so they are written
# plainly here and accents, punctuation and a trailing currency do not matter.
# ``recap`` phrases must be the whole label; the others may also start a
# longer one ("ukupno" reads "UKUPNO II. ZEMLJANI RADOVI") and then only count
# on a line that carries an amount, so a section heading that happens to start
# with "Total" stays a section.
_SUMMARY_KINDS: tuple[str, ...] = ("subtotal", "tax", "grand_total", "recap")

_SUMMARY_WORDS_BY_LANGUAGE: dict[str, dict[str, tuple[str, ...]]] = {
    "en": {
        "subtotal": (
            "total",
            "subtotal",
            "sub total",
            "page total",
            "carried forward",
            "brought forward",
            "carried to collection",
            "total carried to collection",
            "total carried to summary",
            "net total",
            "total excluding vat",
            "total excl vat",
        ),
        "tax": ("vat", "tax", "sales tax", "gst", "hst", "pst", "qst"),
        "grand_total": ("grand total", "total including vat", "total incl vat", "gross total", "contract sum"),
        "recap": ("summary", "collection", "recapitulation", "general summary", "cost summary"),
    },
    "de": {
        "subtotal": ("summe", "zwischensumme", "gesamt", "übertrag", "nettosumme", "summe netto"),
        "tax": ("mwst", "ust", "umsatzsteuer", "mehrwertsteuer"),
        "grand_total": ("gesamtsumme", "gesamtbetrag", "bruttosumme", "summe brutto", "angebotssumme", "endsumme"),
        "recap": ("zusammenstellung", "zusammenfassung"),
    },
    "hr": {
        "subtotal": ("ukupno", "svega", "ukupno bez pdv"),
        "tax": ("pdv",),
        "grand_total": ("sveukupno", "sveukupno s pdv", "ukupno s pdv"),
        "recap": ("rekapitulacija", "rekapitulacija radova", "zbirna rekapitulacija"),
    },
    "sr": {
        "subtotal": ("ukupno", "svega", "укупно", "свега"),
        "tax": ("pdv", "пдв"),
        "grand_total": ("sveukupno", "свеукупно"),
        "recap": ("rekapitulacija", "рекапитулација"),
    },
    "sl": {
        "subtotal": ("skupaj", "vmesni seštevek"),
        "tax": ("ddv",),
        "grand_total": ("skupaj z ddv", "skupna vrednost"),
        "recap": ("rekapitulacija",),
    },
    "pl": {
        "subtotal": ("razem", "suma", "razem netto", "wartość netto"),
        "grand_total": ("ogółem", "razem brutto", "wartość brutto"),
        "recap": ("zestawienie", "podsumowanie"),
    },
    "cs": {
        "subtotal": ("celkem", "mezisoučet", "součet"),
        "tax": ("dph",),
        "grand_total": ("celkem s dph", "celková cena"),
        "recap": ("rekapitulace", "rekapitulace stavby"),
    },
    "sk": {
        "subtotal": ("spolu", "celkom", "medzisúčet"),
        "grand_total": ("spolu s dph", "celkom s dph"),
        "recap": ("rekapitulácia",),
    },
    "hu": {
        "subtotal": ("összesen", "részösszeg", "nettó összesen", "összesen nettó"),
        "tax": ("áfa", "általános forgalmi adó"),
        "grand_total": (
            "mindösszesen",
            "végösszeg",
            "bruttó összesen",
            "összesen bruttó",
            "mindösszesen bruttó",
            "mindösszesen nettó",
        ),
        "recap": ("összesítő", "összesítés", "főösszesítő", "munkanem összesítő"),
    },
    "ro": {
        "subtotal": ("total capitol", "subtotal"),
        "tax": ("tva",),
        "grand_total": ("total general",),
        "recap": ("centralizator", "recapitulatie"),
    },
    "bg": {
        "subtotal": ("общо", "междинна сума"),
        "tax": ("ддс",),
        "grand_total": ("общо с ддс", "всичко"),
        "recap": ("рекапитулация", "обобщение"),
    },
    "el": {
        "subtotal": ("σύνολο", "μερικό σύνολο", "άθροισμα"),
        "tax": ("φπα",),
        "grand_total": ("γενικό σύνολο",),
        "recap": ("ανακεφαλαίωση",),
    },
    "ru": {
        "subtotal": ("итого", "итого по разделу", "всего по разделу"),
        "tax": ("ндс",),
        "grand_total": ("всего", "всего по смете", "итого по смете", "всего с ндс"),
    },
    "uk": {
        "subtotal": ("разом", "всього по розділу", "підсумок"),
        "tax": ("пдв",),
        "grand_total": ("всього", "разом з пдв"),
    },
    "it": {
        "subtotal": ("totale", "subtotale", "totale parziale"),
        "tax": ("iva",),
        "grand_total": ("totale generale", "importo complessivo"),
        "recap": ("riepilogo", "riassunto"),
    },
    "es": {
        "subtotal": ("suma y sigue",),
        "tax": ("igic",),
        "grand_total": ("total general", "total presupuesto"),
        "recap": ("resumen", "resumen de presupuesto"),
    },
    "pt": {
        "subtotal": ("total parcial",),
        "grand_total": ("total geral",),
        "recap": ("resumo",),
    },
    "fr": {
        "subtotal": ("sous total", "total ht"),
        "tax": ("tva",),
        "grand_total": ("total ttc", "montant ttc", "total général"),
        "recap": ("récapitulatif", "récapitulation"),
    },
    "nl": {
        "subtotal": ("totaal", "subtotaal", "totaal excl btw"),
        "tax": ("btw",),
        "grand_total": ("totaal incl btw", "eindtotaal"),
        "recap": ("samenvatting", "recapitulatie"),
    },
    "sv": {
        "subtotal": ("summa", "delsumma", "totalt"),
        "tax": ("moms",),
        "grand_total": ("summa inkl moms", "totalsumma"),
        "recap": ("sammanställning",),
    },
    "no": {
        "subtotal": ("sum", "delsum"),
        "tax": ("mva",),
        "grand_total": ("sum inkl mva", "totalsum"),
        "recap": ("sammendrag", "sammenstilling"),
    },
    "da": {
        "subtotal": ("i alt",),
        "grand_total": ("i alt inkl moms",),
        "recap": ("sammenfatning",),
    },
    "fi": {
        "subtotal": ("yhteensä", "välisumma"),
        "tax": ("alv",),
        "grand_total": ("kokonaissumma", "yhteensä sis alv"),
        "recap": ("yhteenveto",),
    },
    "et": {
        "subtotal": ("kokku", "vahesumma"),
        "tax": ("käibemaks",),
        "grand_total": ("kokku koos käibemaksuga",),
        "recap": ("koond", "kokkuvõte"),
    },
    "tr": {
        "subtotal": ("toplam", "ara toplam"),
        "tax": ("kdv",),
        "grand_total": ("genel toplam",),
        "recap": ("icmal", "özet"),
    },
    "ja": {"subtotal": ("小計", "合計"), "tax": ("消費税",), "grand_total": ("総合計", "総計"), "recap": ("集計",)},
    "zh": {"subtotal": ("小计", "合计"), "tax": ("税金", "增值税"), "grand_total": ("总计",), "recap": ("汇总",)},
    "ko": {"subtotal": ("소계", "합계"), "tax": ("부가세", "부가가치세"), "grand_total": ("총계",), "recap": ("집계",)},
    "ar": {
        "subtotal": ("المجموع", "المجموع الفرعي"),
        "tax": ("ضريبة القيمة المضافة",),
        "grand_total": ("الإجمالي", "المجموع الكلي"),
        "recap": ("ملخص",),
    },
    "he": {
        "subtotal": ('סה"כ', "סיכום ביניים"),
        "tax": ('מע"מ',),
        "grand_total": ('סה"כ כולל מע"מ',),
        "recap": ("ריכוז",),
    },
    "id": {
        "subtotal": ("jumlah", "sub total"),
        "tax": ("ppn",),
        "grand_total": ("jumlah total", "total keseluruhan"),
        "recap": ("rekapitulasi",),
    },
    "vi": {"subtotal": ("cộng",), "tax": ("thuế gtgt",), "grand_total": ("tổng cộng",), "recap": ("tổng hợp",)},
    "th": {"subtotal": ("รวม",), "tax": ("ภาษีมูลค่าเพิ่ม",), "grand_total": ("รวมทั้งสิ้น",)},
    "hi": {"subtotal": ("कुल", "योग"), "tax": ("जीएसटी",), "grand_total": ("कुल योग",)},
    "fa": {"subtotal": ("جمع",), "tax": ("مالیات بر ارزش افزوده",), "grand_total": ("جمع کل",)},
}


def _build_summary_index(table: dict[str, dict[str, tuple[str, ...]]]) -> dict[str, str]:
    """Normalised phrase -> summary kind, across every language."""
    index: dict[str, str] = {}
    for words_by_kind in table.values():
        for kind, phrases in words_by_kind.items():
            for phrase in phrases:
                key = normalise_label(phrase)
                if key:
                    index.setdefault(key, kind)
    return index


_SUMMARY_INDEX: dict[str, str] = _build_summary_index(_SUMMARY_WORDS_BY_LANGUAGE)

# Languages that write the total word at the END of the line: Hungarian names
# the chapter first, "Irtás, föld- és sziklamunka összesen:". Kept to the
# languages that do it, because a trailing "sum" in English is a provisional
# sum, which is work, and matching it everywhere would drop that line.
_TRAILING_SUMMARY_LANGUAGES: tuple[str, ...] = ("hu",)

_TRAILING_SUMMARY_INDEX: dict[str, str] = {
    phrase: kind
    for phrase, kind in _build_summary_index(
        {language: _SUMMARY_WORDS_BY_LANGUAGE[language] for language in _TRAILING_SUMMARY_LANGUAGES}
    ).items()
    if kind != "recap"
}


def summary_label_kind(label: str) -> tuple[str, bool] | None:
    """Say whether a row label reads as a total, tax, grand-total or recap line.

    Args:
        label: The row's description cell.

    Returns:
        ``(kind, whole)`` where ``whole`` is True when the phrase is the entire
        label and False when it only starts it, or ends it in a language of
        :data:`_TRAILING_SUMMARY_LANGUAGES`; ``None`` when no phrase fits.
        A ``recap`` phrase counts only as the whole label.
    """
    normalised = normalise_label(label)
    if not normalised:
        return None
    kind = _SUMMARY_INDEX.get(normalised)
    if kind is not None:
        return kind, True
    words = normalised.split()
    for length in range(len(words) - 1, 0, -1):
        kind = _SUMMARY_INDEX.get(" ".join(words[:length]))
        if kind is not None and kind != "recap":
            return kind, False
    for length in range(len(words) - 1, 0, -1):
        kind = _TRAILING_SUMMARY_INDEX.get(" ".join(words[-length:]))
        if kind is not None:
            return kind, False
    return None


def _is_blank_cell(value: Any) -> bool:
    """A cell that carries nothing: empty, whitespace or a zero."""
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip() or safe_float(value, default=1.0) == 0.0
    return value in (0, 0.0)


def partition_summary_rows(
    rows: list[dict[str, Any]],
    *,
    first_row_number: int = 2,
    row_numbers: list[int] | None = None,
) -> tuple[list[tuple[int, dict[str, Any]]], list[dict[str, Any]]]:
    """Separate a bill's total, tax and recap lines from its sections and work.

    Only a row with no unit, no quantity and no rate can be one of these; a
    priced row is work whatever its label says. Such a row is a summary line
    when its label is a summary phrase in any language of
    :data:`_SUMMARY_WORDS_BY_LANGUAGE` (a phrase that only starts the label
    also needs an amount in the total column), when it carries an amount
    under a recap heading, or when it carries an amount and repeats the
    ordinal of a section already read, which is what a recap page without a
    heading looks like. Every other unpriced row stays a section heading.

    Args:
        rows: Canonical row dicts as the sheet parsers return them.
        first_row_number: The sheet row number of ``rows[0]``, when
            ``row_numbers`` is not given.
        row_numbers: The sheet row of each row, as the Excel reader returns
            them in its metadata.

    Returns:
        ``(kept, summary)``: the rows to import, each with its row number, and
        one report per summary line with its row, ordinal, label, kind and
        amount.
    """
    kept: list[tuple[int, dict[str, Any]]] = []
    summary: list[dict[str, Any]] = []
    section_ordinals: set[str] = set()
    in_recap = False
    numbers = row_numbers if row_numbers is not None and len(row_numbers) == len(rows) else None
    for index, row in enumerate(rows):
        number = numbers[index] if numbers is not None else first_row_number + index
        description = str(row.get("description", "") or "").strip()
        unpriced = (
            not str(row.get("unit", "") or "").strip()
            and _is_blank_cell(row.get("quantity"))
            and _is_blank_cell(row.get("unit_rate"))
        )
        if not description or not unpriced:
            if not unpriced:
                in_recap = False
            kept.append((number, row))
            continue

        amount = safe_float(row.get("total"), default=0.0)
        ordinal = str(row.get("ordinal", "") or "").strip()
        ordinal_key = _label_key(ordinal)
        kind: str | None = None
        matched = summary_label_kind(description)
        if matched is not None and (matched[1] or amount):
            kind = matched[0]
        elif amount and (in_recap or (ordinal_key and ordinal_key in section_ordinals)):
            kind = "recap"

        if kind is None:
            in_recap = False
            if ordinal_key:
                section_ordinals.add(ordinal_key)
            kept.append((number, row))
            continue
        if kind == "recap":
            in_recap = True
        summary.append(
            {
                "row": number,
                "ordinal": ordinal,
                "description": description[:200],
                "kind": kind,
                "amount": amount,
            }
        )
    return kept, summary


_SUMMARY_KIND_WORDS: dict[str, str] = {
    "subtotal": "subtotal",
    "tax": "tax",
    "grand_total": "grand total",
    "recap": "recap",
}


def summary_row_warning(report: dict[str, Any]) -> dict[str, Any]:
    """The import warning that tells the user a summary line was left out."""
    return {
        "row": report["row"],
        "ordinal": report["ordinal"],
        "severity": "info",
        "code": "summary_row_skipped",
        "kind": report["kind"],
        "amount": report["amount"],
        # The row number travels in ``row``; the import dialog prints it in
        # front of the message, so the message does not repeat it.
        "label": report["description"][:80],
        "message": (
            f"'{report['description'][:80]}' reads as a "
            f"{_SUMMARY_KIND_WORDS[report['kind']]} line and was not imported as a position."
        ),
    }


def _split_metadata(row: dict[str, Any], language: str | None) -> dict[str, Any]:
    """The material and labour halves of a line, keyed for where they are read."""
    if not any(key in row for key in _SPLIT_COLUMNS):
        return {}
    halves = {key: safe_float(row.get(key), default=0.0) for key in _SPLIT_COLUMNS}
    if language == "hu":
        return {
            "hu": {
                "profile": "flat",
                "material_unit_rate": halves["material_rate"],
                "fee_unit_rate": halves["labour_rate"],
                "material_total": halves["material_total"],
                "fee_total": halves["labour_total"],
            }
        }
    return {
        "rate_split": {
            "material_unit_rate": halves["material_rate"],
            "labour_unit_rate": halves["labour_rate"],
            "material_total": halves["material_total"],
            "labour_total": halves["labour_total"],
        }
    }


# Contingency lines, matched on :func:`normalise_label` at the start of the
# label. A bill writes its reserve as a line with a name and an amount and no
# unit, quantity or rate, "Tartalékkeret 5%", which is exactly the shape of a
# section heading, so it imported as an empty section and the reserve was
# lost from the total. Such a line is money the client budgets, so it comes in
# as a lump sum carrying its amount.
_CONTINGENCY_PHRASES: frozenset[str] = frozenset(
    normalise_label(phrase)
    for phrase in (
        # English
        "contingency",
        "contingencies",
        "contingency sum",
        "contingency allowance",
        # Hungarian
        "tartalékkeret",
        "tartalék",
        "előre nem látható költségek",
    )
)


def _contingency_amount(row: dict[str, Any], description: str) -> float | None:
    """The amount of a contingency line written without a unit, quantity or rate.

    Returns ``None`` for any other row, and for a contingency line that is
    priced like work already or carries no amount.
    """
    if str(row.get("unit", "") or "").strip():
        return None
    if not (_is_blank_cell(row.get("quantity")) and _is_blank_cell(row.get("unit_rate"))):
        return None
    words = normalise_label(description).split()
    if not any(" ".join(words[:length]) in _CONTINGENCY_PHRASES for length in range(1, len(words) + 1)):
        return None
    amount = safe_float(row.get("total"), default=0.0)
    return amount if amount > 0 else None


_IMPORT_MAX_QUANTITY = 1e9
_IMPORT_MAX_UNIT_RATE = 1e8


def _rows_to_positions(
    rows: list[dict[str, Any]],
    *,
    source: str = "excel_import",
    row_numbers: list[int] | None = None,
    header_language: str | None = None,
) -> ImportedBOQ:
    """Convert canonical rows into :class:`ImportedPosition` objects.

    Carries the sanity bounds + section-row + summary-row detection that
    the legacy inline parser used. Per-row errors are collected on the
    returned :class:`ImportedBOQ` rather than raised so the dispatcher
    can return them as a structured list.

    A line read from material and labour columns keeps the two halves in its
    metadata. A Hungarian bill keeps them under ``hu`` with the keys the
    workbook profile writes, which is where the Hungarian material and fee
    rule reads them; any other bill keeps them under ``rate_split``.
    """
    result = ImportedBOQ(source_format="csv-or-xlsx")
    auto_ordinal = 1

    # Pre-compute a median unit rate across the file so we can warn on
    # any single position that's >10× above (likely a tampered export).
    rate_samples = sorted(v for v in (safe_float(r.get("unit_rate"), default=0.0) for r in rows) if v > 0)
    median_rate = rate_samples[len(rate_samples) // 2] if rate_samples else 0.0

    kept_rows, summary_rows = partition_summary_rows(rows, row_numbers=row_numbers)
    result.skipped += len(summary_rows)
    result.warnings.extend(summary_row_warning(report) for report in summary_rows)
    if summary_rows:
        result.metadata["summary_rows"] = summary_rows

    # Numbers the file writes itself. A generated ordinal steps over them: a
    # bill that numbers its lines but not its chapter headings ("Ssz." blank
    # on the heading row) used to give the first heading "1" beside line 1,
    # and duplicate ordinals fail validation.
    explicit_ordinals = {str(row.get("ordinal", "")).strip() for _, row in kept_rows} - {""}

    # A national bill that restarts its numbering at every part (Bulgarian
    # KCC-style bills do) yields the same explicit ordinal once per section.
    # ``boq_quality.no_duplicate_ordinals`` is a hard error, so a repeat is
    # qualified with the running section index and the document's own number
    # is preserved in ``metadata.source_ordinal``.
    seen_ordinals: set[str] = set()
    section_seq = 0

    for row_idx, row in kept_rows:
        try:
            description = str(row.get("description", "")).strip()
            if not description:
                result.skipped += 1
                continue

            desc_lower = description.lower()
            if desc_lower in _TOTAL_ROW_DESCRIPTIONS:
                result.skipped += 1
                continue
            if desc_lower.startswith("subtotal:") or desc_lower.startswith("zwischensumme:"):
                result.skipped += 1
                continue

            # A national bill that re-prints its column headings at the top of
            # every section yields a data row made only of header words; it is
            # never a real position.
            if sum(
                1 for key in ("description", "unit", "quantity", "unit_rate")
                if _match_column(str(row.get(key) or ""))
            ) >= 2:
                result.skipped += 1
                continue

            ordinal = str(row.get("ordinal", "")).strip()
            if not ordinal:
                while str(auto_ordinal) in explicit_ordinals:
                    auto_ordinal += 1
                ordinal = str(auto_ordinal)
            auto_ordinal += 1

            source_ordinal = ""
            if ordinal in seen_ordinals:
                qualified = f"{section_seq}.{ordinal}"
                suffix = 2
                while qualified in seen_ordinals:
                    qualified = f"{section_seq}.{ordinal}.{suffix}"
                    suffix += 1
                source_ordinal, ordinal = ordinal, qualified
            seen_ordinals.add(ordinal)

            # Round-trip identity (GitHub #360): the dedicated "Position ID"
            # column an export stamped. Blank -> new row; a value belonging to
            # the target BOQ -> update in place (resolved downstream by the
            # diff against the BOQ's current ids).
            position_id = normalise_id(row.get("position_id"))

            unit_raw = str(row.get("unit", "")).strip()
            quantity_raw = row.get("quantity")
            unit_rate_raw = row.get("unit_rate")
            contingency = _contingency_amount(row, description)
            if contingency is not None:
                unit_raw, quantity_raw, unit_rate_raw = "lsum", 1.0, contingency
            quantity, q_err = parse_numeric_cell(quantity_raw)
            unit_rate, r_err = parse_numeric_cell(unit_rate_raw)
            if q_err is not None:
                result.errors.append(
                    {
                        "row": row_idx,
                        "ordinal": ordinal,
                        "error": f"Invalid quantity at row {row_idx}: {q_err}",
                    }
                )
                continue
            if r_err is not None:
                result.errors.append(
                    {
                        "row": row_idx,
                        "ordinal": ordinal,
                        "error": f"Invalid unit_rate at row {row_idx}: {r_err}",
                    }
                )
                continue
            assert quantity is not None
            assert unit_rate is not None

            # Section detection: a row with a description but no unit /
            # quantity / rate is a section header from our own exporter.
            is_section_row = (
                not unit_raw and (quantity_raw in (None, "", 0, 0.0)) and (unit_rate_raw in (None, "", 0, 0.0))
            )
            if is_section_row:
                section_seq += 1
                section_meta: dict[str, Any] = {
                    "import_row_index": row_idx,
                    "section_header": True,
                }
                if source_ordinal:
                    section_meta["source_ordinal"] = source_ordinal
                result.positions.append(
                    ImportedPosition(
                        description=description,
                        ordinal=ordinal,
                        unit="section",
                        quantity=0.0,
                        unit_rate=0.0,
                        classification={},
                        source=source,
                        metadata=section_meta,
                        is_section=True,
                        position_id=position_id,
                    )
                )
                continue

            # A trailing prime/foot mark is notation for the linear unit
            # itself ("м'", "ft'"), not part of the symbol, and the unit
            # field validator rejects the apostrophe.
            unit = (unit_raw or "pcs").rstrip("'′’") or "pcs"

            # Range guards - reject obvious tamper / typo errors.
            if not (0 <= quantity <= _IMPORT_MAX_QUANTITY):
                result.errors.append(
                    {
                        "row": row_idx,
                        "ordinal": ordinal,
                        "error": f"Quantity out of range: {quantity}",
                    }
                )
                continue
            if not (0 <= unit_rate <= _IMPORT_MAX_UNIT_RATE):
                result.errors.append(
                    {
                        "row": row_idx,
                        "ordinal": ordinal,
                        "error": f"Unit rate out of range: {unit_rate}",
                    }
                )
                continue

            # Soft warnings. A lump sum prices a whole piece of work, so its
            # rate says nothing next to the per-metre rates around it.
            if median_rate > 0 and unit_rate > median_rate * 10 and not is_lump_sum_unit(unit):
                result.warnings.append(
                    {
                        "row": row_idx,
                        "ordinal": ordinal,
                        "severity": "warning",
                        "message": (
                            f"Unit rate {unit_rate:.2f} is >10× the file median "
                            f"({median_rate:.2f}) - possible typo or tampered export."
                        ),
                    }
                )
            if quantity == 0:
                result.warnings.append(
                    {
                        "row": row_idx,
                        "ordinal": ordinal,
                        "severity": "info",
                        "message": "Quantity is zero - position imported but contributes no cost.",
                    }
                )
            if unit_rate == 0:
                result.warnings.append(
                    {
                        "row": row_idx,
                        "ordinal": ordinal,
                        "severity": "info",
                        "message": "Unit rate is zero - position imported without a rate.",
                    }
                )

            # Heuristic classification (Epics I9 + I10).
            class_value = str(row.get("classification", "")).strip()
            classification = _infer_classification(class_value, description)
            if header_language == "hu" and "code" in classification:
                # The Hungarian rules read the item code from ``tetelrend``,
                # where the workbook profile writes it. A flat bill's code is
                # carried there too, so those rules judge the code the line
                # has rather than report it as having none.
                classification["tetelrend"] = classification["code"]

            metadata: dict[str, Any] = {"import_row_index": row_idx}
            if source_ordinal:
                metadata["source_ordinal"] = source_ordinal
            if contingency is not None:
                metadata["contingency"] = True
            split = _split_metadata(row, header_language)
            if split:
                metadata.update(split)

            result.positions.append(
                ImportedPosition(
                    description=description,
                    ordinal=ordinal,
                    unit=unit,
                    quantity=quantity,
                    unit_rate=unit_rate,
                    classification=classification,
                    source=source,
                    metadata=metadata,
                    position_id=position_id,
                )
            )

        except Exception as exc:  # noqa: BLE001 - caller surfaces row #
            result.errors.append({"row": row_idx, "ordinal": "", "error": str(exc)})
            logger.warning("Excel/CSV row %d error: %s", row_idx, exc)

    return result


class ExcelImporter:
    """Generic Excel (.xlsx) / CSV importer with NRM + MasterFormat heuristics."""

    format_id: ClassVar[str] = "excel"
    extensions: ClassVar[tuple[str, ...]] = (".xlsx", ".csv")
    display_name: ClassVar[str] = "Excel / CSV BOQ"
    rule_packs: ClassVar[tuple[str, ...]] = ("boq_quality",)

    @classmethod
    def detect(cls, head_bytes: bytes, filename: str) -> bool:
        """Detect by magic bytes (xlsx zip header / CSV text) + extension."""
        if not head_bytes:
            return False
        name = filename.lower()
        if not any(name.endswith(ext) for ext in cls.extensions):
            return False
        fmt = _detect_file_format(head_bytes[:4096])
        if name.endswith(".xlsx"):
            return fmt == "xlsx"
        if name.endswith(".csv"):
            return fmt == "csv"
        return False

    @classmethod
    async def parse(cls, content: bytes, *, locale: str = "en") -> ImportedBOQ:
        """Parse an .xlsx or .csv BOQ into :class:`ImportedBOQ`."""
        if not content:
            raise ImporterParseError("Spreadsheet upload is empty")

        fmt = _detect_file_format(content[:4096])

        # Hungarian bills are not tables with a header row, so the alias mapper
        # below reads nothing out of them: the building shape spreads its item
        # code across nine columns on seventeen sheets, and both shapes carry
        # two unit prices per line rather than one. The profile answers only for
        # a workbook it recognises and hands everything else straight back, and
        # it has to be asked here rather than in ``detect`` because an xlsx is a
        # zip whose first four kilobytes say nothing about its contents.
        if fmt == "xlsx":
            hungarian = parse_hungarian_workbook(content)
            if hungarian is not None and hungarian.positions:
                return hungarian

        import_meta: dict[str, Any] = {}
        try:
            if fmt == "xlsx":
                rows, import_meta = _parse_rows_from_excel(content)
                source_format = "xlsx"
            elif fmt == "csv":
                rows, import_meta = _parse_csv(content)
                source_format = "csv"
            else:
                raise ImporterParseError(f"Unsupported spreadsheet format: detected {fmt!r}")
        except ImporterParseError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ImporterParseError(f"Could not parse spreadsheet: {exc}") from exc

        if not rows:
            raise ImporterParseError("No data rows found. Check that the header row names the columns.")

        language = import_meta.get("header_language")
        result = _rows_to_positions(rows, row_numbers=import_meta.get("row_numbers"), header_language=language)
        result.source_format = source_format
        result.metadata = {
            **result.metadata,
            "original_columns": import_meta.get("original_columns", []),
            "column_mapping": import_meta.get("column_mapping", {}),
            "sheet_names": import_meta.get("sheet_names", []),
            "total_rows_seen": len(rows),
        }
        for key in ("header_language", "item_sheet", "encoding", "delimiter"):
            if import_meta.get(key):
                result.metadata[key] = import_meta[key]
        return result
