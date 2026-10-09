"""Offline Bulgarian KCC replay against a CSV export, using the ERP matcher.

The export has no DB metadata or pattern priors; this is not a PostgreSQL
integration test. Only operator-origin prices are automatically applied.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from app.modules.cost_match.bulgarian import identity_text, is_structural, parse_work, quotation_eligible
from app.modules.cost_match.matcher import (
    _GENERIC_TERMS,
    _PROCESS_CONCEPTS,
    Candidate,
    best_match,
    canonical_tokens,
    normalize_unit,
    suggestion_rate,
    unit_rate_factor,
    unit_scale,
    work_rate_factor,
)
from app.modules.cost_match.retrieval_bg import literal_description, retrieval_plan
from app.modules.cost_match.work_catalog import evidence_origin, work_metadata
from app.modules.cost_match.work_context import effective_work, is_work_fragment

OPERATOR_SOURCES = {'operator_pricelist', 'manual_entry', 'user_upload', 'client_directive', 'custom', 'estimate_confirmed'}


def decimal(value):
    try:
        result = Decimal(str(value).replace(' ', '').replace(',', '.'))
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError, TypeError):
        return None


def content(text):
    from app.modules.cost_match.work_catalog import notation

    values, remainder = notation(text, parse_work(text).object)
    terms = {t for t in canonical_tokens(identity_text(remainder))
            if t not in _GENERIC_TERMS and t not in _PROCESS_CONCEPTS
            and len(t) > 1 and any(c.isalpha() for c in t)}
    terms.update(f'work:{name}:{value}' for name, items in values for value in items)
    return terms


def quotation_key(code, text, unit, rate, currency):
    return (code, literal_description(text), normalize_unit(unit), unit_scale(unit),
            decimal(rate), (currency or '').strip().upper())


def seed_provenance(corpus):
    index = {}
    conflicts = set()
    for path in sorted(corpus.parent.parent.glob('*operator*.json')):
        try:
            records = json.loads(path.read_text(encoding='utf-8-sig'))
        except (OSError, ValueError) as error:
            print(f'Provenance not loaded from {path.name}: {error}', flush=True)
            continue
        if not isinstance(records, list):
            continue
        for row in records:
            if not isinstance(row, dict):
                continue
            money = row.get('money')
            if not isinstance(money, dict) or not evidence_origin(row):
                continue
            key = quotation_key(row.get('code') or row.get('id'), row.get('desc') or row.get('name') or '',
                                row.get('unit'), money.get('amount'), money.get('currency'))
            info = {'evidence_origin': evidence_origin(row), 'vat_included': money.get('vatIncluded')}
            if key in index and index[key] != info:
                conflicts.add(key)
            else:
                index[key] = info
    return {key: value for key, value in index.items() if key not in conflicts}


def run(corpus, input_path, output_path, limit):
    if input_path.resolve() == output_path.resolve():
        raise ValueError('Output must not overwrite the original input workbook.')
    candidates = []
    postings = defaultdict(set)
    exact = defaultdict(set)
    literal_exact = defaultdict(set)
    signature_postings = defaultdict(set)
    price_spreads = defaultdict(list)
    objects = []
    tokens = []
    scanned = 0
    provenance = seed_provenance(corpus)
    origin_counts = Counter()
    with corpus.open(encoding='utf-8-sig', newline='') as handle:
        for row in csv.DictReader(handle):
            scanned += 1
            if str(row.get('is_active', 'true')).lower() not in {'true', '1'}:
                continue
            rate = decimal(row.get('rate'))
            currency = row.get('currency', '').strip().upper()
            if rate is None or rate <= 0 or currency not in {'EUR', 'BGN'}:
                continue
            if currency == 'BGN':
                rate /= Decimal('1.95583')
            try:
                metadata = json.loads(row.get('classification') or '{}')
            except (TypeError, ValueError):
                metadata = {}
            if not isinstance(metadata, dict):
                metadata = {}
            seed_info = provenance.get(quotation_key(row.get('code'), row.get('description', ''), row.get('unit'),
                                                     row.get('rate'), row.get('currency')), {})
            origin = evidence_origin(metadata) or row.get('source_boq') or seed_info.get('evidence_origin', '')
            origin_counts['documented' if origin else 'unknown'] += 1
            text, unit = row.get('description', ''), row.get('unit', '')
            price_spreads[literal_description(text)].append((
                rate, unit, row.get('source', ''),
                (row.get('price_as_of') or (row.get('created_at') if row.get('source') == 'estimate_confirmed' else '') or '')[:10],
            ))
            index = len(candidates)
            signature = work_metadata(text)
            candidates.append(Candidate(str(index), text, unit, {
                'code': row['code'], 'source': row['source'],
                'unit_rate': str(rate), 'currency': 'EUR',
                'quotation_rate': row.get('rate'), 'quotation_currency': row.get('currency'),
                'canonical_work': signature,
                'auto_pricing_eligible': row.get('source') in OPERATOR_SOURCES,
                'provenance_status': 'approved' if row.get('source') in OPERATOR_SOURCES else 'pending_review',
                'price_as_of': row.get('price_as_of') or '',
                'created_at': row.get('created_at') or '',
                'evidence_origin': origin,
                'vat_included': metadata.get('vat_included', seed_info.get('vat_included')),
            }))
            exact[identity_text(text)].add(index)
            literal_exact[literal_description(text)].add(index)
            for kind, values in signature['work']['concept_ids'].items():
                values = [values] if kind in {'object', 'scope'} else values
                for value in values:
                    signature_postings[kind, value].add(index)
            terms = content(text)
            tokens.append(terms)
            objects.append(parse_work(text).object)
            for term in terms:
                postings[term].add(index)
    print(f'CSV rows scanned: {scanned}; active price evidence records: {len(candidates)}', flush=True)
    print(f'Price origins: {dict(origin_counts)}', flush=True)

    workbook = load_workbook(input_path)
    sheet = workbook.active
    cached_values = load_workbook(input_path, data_only=True, read_only=True)
    cached_sheet = cached_values[sheet.title]
    audit = workbook.create_sheet('BG_MATCH_AUDIT')
    audit.append(['Excel row', 'Position', 'Query', 'Unit', 'Decision', 'EUR unit rate',
                  'Source code', 'Source', 'Source description', 'Confidence',
                  'Reasons', 'Object', 'Operations', 'Scope', 'Specifications', 'Alternatives',
                  'KCC work identity', 'Quote work identity', 'Inherited work context',
                   'Effective work description', 'Retrieval plan', 'Candidates before limit',
                   'Equivalent price pool'])
    sheet.cell(1, 5, 'Единична цена EUR')
    sheet.cell(1, 6, 'Обща цена EUR')
    counts = Counter()
    cache = {}
    retrieval_counts = {}
    parent_work = ''
    for row_no in range(2, sheet.max_row + 1):
        position = str(sheet.cell(row_no, 1).value or '')
        description = str(sheet.cell(row_no, 2).value or '').strip()
        unit = str(sheet.cell(row_no, 3).value or '').strip()
        quantity_cell = sheet.cell(row_no, 4)
        quantity = decimal(cached_sheet.cell(row_no, 4).value
                           if quantity_cell.data_type == 'f' else quantity_cell.value)
        if description and (not unit or unit.lower() == 'section'
                            or quantity_cell.value in (None, '')):
            parent_work = description
            counts['structural_or_nonwork'] += 1
            continue
        if not description or not unit or unit.lower() == 'section' or is_structural(description):
            parent_work = description
            if description:
                counts['structural_or_nonwork'] += 1
            continue
        inherited = [parent_work] if parent_work and is_work_fragment(description) else []
        effective = effective_work(description, inherited)
        if not is_work_fragment(description):
            parent_work = ''
        if sheet.cell(row_no, 5).value not in (None, ''):
            counts['existing_rate_preserved'] += 1
            continue
        if quantity is None:
            counts['review'] += 1
            audit.append([row_no, position, description, unit, 'review', None,
                          None, None, None, None, 'quantity_unavailable'])
            continue
        key = (effective, unit, description)
        if key not in cache:
            terms = content(effective)
            plan = retrieval_plan(effective, description)
            possible = set(range(len(candidates))) if not plan.blocked_reason else set()
            for kind, values in plan.id_requirements().items():
                values = [values] if kind == 'object' else values
                for value in values:
                    possible.intersection_update(signature_postings[kind, value])
            if not plan.groups:
                possible = set(literal_exact[plan.exact_text]) if plan.exact_text is not None else set()
            possible = {i for i in possible if quotation_eligible(description, candidates[i].text)}
            if plan.candidate_code:
                possible = {i for i in possible if candidates[i].payload['code'] == plan.candidate_code}
            retrieval_counts[key] = len(possible)
            obj = parse_work(effective).object
            ordered = sorted(possible, key=lambda i: (
                i not in exact.get(identity_text(effective), ()),
                objects[i] != obj,
                unit_rate_factor(candidates[i].unit, unit) is None,
                -len(terms & tokens[i]),
                candidates[i].payload['code'],
            ))
            cache[key] = best_match(effective, [candidates[i] for i in ordered], query_unit=unit, locale='bg', raw_query=description)
        outcome = cache[key]
        frame = parse_work(effective)
        selected = outcome.candidate
        score = outcome.score
        reason = list(score.reasons) if score else ['no_candidates']
        applied = None
        decision = 'review' if selected else 'unmatched'
        if selected:
            reference_price = bool(score.factors.get('reference_price')) or (selected.payload or {}).get('auto_pricing_eligible') is False or any(
                c.ref in outcome.pool_refs
                and (c.payload or {}).get('auto_pricing_eligible') is False
                for c, _ in outcome.scored_all
            )
            if reference_price:
                reason.append('reference_price')
            factor = work_rate_factor(selected.unit, unit, query_text=effective, candidate_text=selected.text)
            rate = suggestion_rate(selected)
            priced_candidate = outcome.median_candidate or selected
            operator_price = (priced_candidate.payload or {}).get('source') in OPERATOR_SOURCES
            peers = [(rate * f, source, day)
                     for rate, source_unit, source, day in price_spreads[literal_description(priced_candidate.text)]
                     if (source in OPERATOR_SOURCES) == operator_price
                     and (f := unit_rate_factor(source_unit, priced_candidate.unit)) is not None]
            confirmed = [peer for peer in peers if peer[1] == 'estimate_confirmed']
            if confirmed:
                peers = confirmed
            if peers and all(day for _, _, day in peers):
                newest = max(day for _, _, day in peers)
                peers = [peer for peer in peers if peer[2] == newest]
            spread = [rate for rate, _, _ in peers]
            disagreement = spread and max(spread) > min(spread) * Decimal('1.5')
            if disagreement:
                reason.append('corpus_price_disagreement')
            if (outcome.is_confident and not reference_price and not disagreement and not outcome.pool_diverged
                    and rate is not None and factor is not None and quantity is not None
                    and quantity >= 0):
                applied = (outcome.median_rate if outcome.median_rate is not None else rate * factor).quantize(Decimal('0.0001'))
                decision = 'exact' if score.factors.get('exact') else 'compatible'
                sheet.cell(row_no, 5, float(applied))
                sheet.cell(row_no, 6, float((applied * quantity).quantize(Decimal('0.01'))))
        counts[decision] += 1
        payload = selected.payload if selected else {}
        alternatives = [{
            'code': c.payload['code'], 'description': c.text, 'confidence': s.confidence,
            'unit': c.unit, 'EUR_source_rate': c.payload['unit_rate'], 'reasons': s.reasons,
        } for c, s in outcome.alternatives]
        audit.append([row_no, position, description, unit, decision,
                      float(applied) if applied is not None else None,
                      payload.get('code'), payload.get('source'), selected.text if selected else None,
                      score.confidence if score else 0, '; '.join(reason), frame.object,
                      ', '.join(sorted(frame.operations)), frame.scope,
                      json.dumps({k: sorted(v) for k, v in frame.specs}, ensure_ascii=False),
                      json.dumps(alternatives, ensure_ascii=False),
                      json.dumps(work_metadata(effective), ensure_ascii=False),
                      json.dumps(payload.get('canonical_work'), ensure_ascii=False),
                       json.dumps(inherited, ensure_ascii=False), effective,
                       json.dumps(retrieval_plan(effective, description).as_dict(), ensure_ascii=False),
                       retrieval_counts[key], json.dumps({
                           'median_EUR': None if outcome.median_rate is None else str(outcome.median_rate),
                           'min': str(outcome.pool_min), 'max': str(outcome.pool_max),
                           'diverged': outcome.pool_diverged,
                           'observations': [dict(ref=c.ref, description=c.text, unit=c.unit,
                                                in_pool=c.ref in outcome.pool_refs, **c.payload)
                                            for c, _ in outcome.scored_all
                                            if c.ref in outcome.pool_refs],
                       }, ensure_ascii=False)])
    audit.freeze_panes = 'A2'
    audit.auto_filter.ref = audit.dimensions
    cached_values.close()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(json.dumps(dict(counts), ensure_ascii=False), flush=True)
    print(str(output_path.resolve()), flush=True)
    return counts


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--corpus', type=Path, required=True)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--limit', type=int, default=100)
    args = parser.parse_args()
    if args.limit <= 0:
        parser.error('--limit must be positive')
    run(args.corpus, args.input, args.output, args.limit)
