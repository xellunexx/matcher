"""Offline Bulgarian KCC replay against a CSV export, using the ERP matcher.

The export has no DB metadata or pattern priors; this is not a PostgreSQL
integration test. Only operator-origin prices are automatically applied.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from app.modules.cost_match.bulgarian import identity_text, is_structural, parse_work
from app.modules.cost_match.work_catalog import work_metadata
from app.modules.cost_match.work_context import effective_work, is_work_fragment
from app.modules.cost_match.matcher import (
    _GENERIC_TERMS, _PROCESS_CONCEPTS, Candidate, best_match,
    canonical_tokens, suggestion_rate, unit_rate_factor,
)


OPERATOR_SOURCES = {'operator_pricelist', 'manual_entry', 'user_upload', 'client_directive', 'custom'}


def decimal(value):
    try:
        result = Decimal(str(value).replace(' ', '').replace(',', '.'))
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError, TypeError):
        return None


def content(text):
    from app.modules.cost_match.work_catalog import notation

    values, remainder = notation(text, parse_work(text).object)
    terms = {t for t in canonical_tokens(identity_text(remainder))
            if t not in _GENERIC_TERMS and t not in _PROCESS_CONCEPTS
            and len(t) > 1 and any(c.isalpha() for c in t)}
    terms.update(f'work:{name}:{value}' for name, items in values for value in items)
    return terms


def run(corpus, input_path, output_path, limit):
    if input_path.resolve() == output_path.resolve():
        raise ValueError('Output must not overwrite the original input workbook.')
    candidates = []
    postings = defaultdict(set)
    exact = defaultdict(set)
    price_spreads = defaultdict(list)
    objects = []
    tokens = []
    scanned = 0
    with corpus.open(encoding='utf-8-sig', newline='') as handle:
        for row in csv.DictReader(handle):
            scanned += 1
            if str(row.get('is_active', 'true')).lower() not in {'true', '1'}:
                continue
            rate = decimal(row.get('rate'))
            currency = row.get('currency', '').strip().upper()
            if rate is None or rate <= 0 or currency not in {'EUR', 'BGN'}:
                continue
            if currency == 'BGN':
                rate /= Decimal('1.95583')
            text, unit = row.get('description', ''), row.get('unit', '')
            price_spreads[text.strip().lower()].append((rate, unit))
            if row.get('source') not in OPERATOR_SOURCES:
                continue
            index = len(candidates)
            candidates.append(Candidate(str(index), text, unit, {
                'code': row['code'], 'source': row['source'],
                'unit_rate': str(rate), 'currency': 'EUR',
                'canonical_work': work_metadata(text),
            }))
            exact[identity_text(text)].add(index)
            terms = content(text)
            tokens.append(terms)
            objects.append(parse_work(text).object)
            for term in terms:
                postings[term].add(index)
    print(f'CSV rows scanned: {scanned}; operator price records: {len(candidates)}', flush=True)

    workbook = load_workbook(input_path)
    sheet = workbook.active
    audit = workbook.create_sheet('BG_MATCH_AUDIT')
    audit.append(['Excel row', 'Position', 'Query', 'Unit', 'Decision', 'EUR unit rate',
                  'Source code', 'Source', 'Source description', 'Confidence',
                  'Reasons', 'Object', 'Operations', 'Scope', 'Specifications', 'Alternatives',
                  'KCC work identity', 'Quote work identity', 'Inherited work context',
                  'Effective work description'])
    sheet.cell(1, 5, 'Единична цена EUR')
    sheet.cell(1, 6, 'Обща цена EUR')
    counts = Counter()
    cache = {}
    parent_work = ''
    for row_no in range(2, sheet.max_row + 1):
        position = str(sheet.cell(row_no, 1).value or '')
        description = str(sheet.cell(row_no, 2).value or '').strip()
        unit = str(sheet.cell(row_no, 3).value or '').strip()
        quantity = decimal(sheet.cell(row_no, 4).value)
        if description and (not unit or unit.lower() == 'section' or quantity is None):
            parent_work = description
            counts['structural_or_nonwork'] += 1
            continue
        if not description or not unit or unit.lower() == 'section' or is_structural(description):
            parent_work = description
            if description:
                counts['structural_or_nonwork'] += 1
            continue
        inherited = [parent_work] if parent_work and is_work_fragment(description) else []
        effective = effective_work(description, inherited)
        if not is_work_fragment(description):
            parent_work = ''
        if sheet.cell(row_no, 5).value not in (None, ''):
            counts['existing_rate_preserved'] += 1
            continue
        key = (effective, unit)
        if key not in cache:
            terms = content(effective)
            possible = set(exact.get(identity_text(effective), ()))
            for term in terms:
                possible.update(postings.get(term, ()))
            obj = parse_work(effective).object
            ordered = sorted(possible, key=lambda i: (
                i not in exact.get(identity_text(effective), ()),
                objects[i] != obj,
                unit_rate_factor(candidates[i].unit, unit) is None,
                -len(terms & tokens[i]),
                candidates[i].payload['code'],
            ))[:limit]
            cache[key] = best_match(effective, [candidates[i] for i in ordered], query_unit=unit, locale='bg')
        outcome = cache[key]
        frame = parse_work(effective)
        selected = outcome.candidate
        score = outcome.score
        reason = list(score.reasons) if score else ['no_candidates']
        applied = None
        decision = 'review' if selected else 'unmatched'
        if selected:
            factor = unit_rate_factor(selected.unit, unit)
            rate = suggestion_rate(selected)
            spread = [rate * f for rate, source_unit in price_spreads[selected.text.strip().lower()]
                      if (f := unit_rate_factor(source_unit, selected.unit)) is not None]
            disagreement = spread and max(spread) > min(spread) * Decimal('1.5')
            if disagreement:
                reason.append('corpus_price_disagreement')
            if (outcome.is_confident and not disagreement and not outcome.pool_diverged
                    and rate is not None and factor is not None and quantity is not None
                    and quantity >= 0):
                applied = (rate * factor).quantize(Decimal('0.0001'))
                decision = 'exact' if score.factors.get('exact') else 'compatible'
                sheet.cell(row_no, 5, float(applied))
                sheet.cell(row_no, 6, float((applied * quantity).quantize(Decimal('0.01'))))
        counts[decision] += 1
        payload = selected.payload if selected else {}
        alternatives = [{
            'code': c.payload['code'], 'description': c.text, 'confidence': s.confidence,
            'unit': c.unit, 'EUR_source_rate': c.payload['unit_rate'], 'reasons': s.reasons,
        } for c, s in outcome.alternatives]
        audit.append([row_no, position, description, unit, decision,
                      float(applied) if applied is not None else None,
                      payload.get('code'), payload.get('source'), selected.text if selected else None,
                      score.confidence if score else 0, '; '.join(reason), frame.object,
                      ', '.join(sorted(frame.operations)), frame.scope,
                      json.dumps({k: sorted(v) for k, v in frame.specs}, ensure_ascii=False),
                      json.dumps(alternatives, ensure_ascii=False),
                      json.dumps(work_metadata(effective), ensure_ascii=False),
                      json.dumps(payload.get('canonical_work'), ensure_ascii=False),
                      json.dumps(inherited, ensure_ascii=False), effective])
    audit.freeze_panes = 'A2'
    audit.auto_filter.ref = audit.dimensions
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(json.dumps(dict(counts), ensure_ascii=False), flush=True)
    print(str(output_path.resolve()), flush=True)
    return counts


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--corpus', type=Path, required=True)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--limit', type=int, default=100)
    args = parser.parse_args()
    if args.limit <= 0:
        parser.error('--limit must be positive')
    run(args.corpus, args.input, args.output, args.limit)
