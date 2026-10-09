# -*- coding: utf-8 -*-
import sys, io, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

def load(p):
    wb = openpyxl.load_workbook(p, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows

def analyze(tag, p):
    rows = load(p)
    print(f'===== {tag}: {p} ({len(rows)} rows)')
    # header row index 1 (0-based 1)
    hdr = rows[1]
    cols = {name: i for i, name in enumerate(hdr) if name}
    print('cols:', {k: cols[k] for k in cols})
    # find columns
    i_desc = cols.get('Description'); i_rate = cols.get('Unit Rate')
    i_total = cols.get('Total')
    priced = unpriced = 0
    rates = []
    sample_priced = []
    for r in rows[2:]:
        d = r[i_desc] if i_desc is not None and i_desc < len(r) else None
        rt = r[i_rate] if i_rate is not None and i_rate < len(r) else None
        if d is None or str(d).strip() == '':
            continue
        try:
            v = float(rt) if rt is not None else None
        except (TypeError, ValueError):
            v = None
        if v is None:
            unpriced += 1
        else:
            priced += 1
            rates.append(v)
            if len(sample_priced) < 12:
                sample_priced.append((str(d)[:60], v))
    print(f'priced={priced} unpriced={unpriced}')
    dist = collections.Counter(round(x, 2) for x in rates)
    print('top repeated rates:', dist.most_common(12))
    for d, v in sample_priced:
        print(f'  {v:>10.2f}  {d}')
    return rows, cols

b_rows, b_cols = analyze('BEFORE', r'C:\Users\ochak\Downloads\kcc2 (3).xlsx')
a_rows, a_cols = analyze('AFTER', r'C:\Users\ochak\Downloads\kcc2 (4).xlsx')

# diff by position no + description
def keymap(rows, cols):
    i_pos = cols.get('Pos.'); i_desc = cols.get('Description'); i_rate = cols.get('Unit Rate')
    m = {}
    for r in rows[2:]:
        if i_desc is None or r[i_desc] is None: continue
        k = (r[i_pos], str(r[i_desc])[:80])
        m[k] = r[i_rate]
    return m

bm = keymap(b_rows, b_cols); am = keymap(a_rows, a_cols)
changed = []
for k, bv in bm.items():
    av = am.get(k)
    if bv != av:
        changed.append((k, bv, av))
print(f'===== DIFF: {len(changed)} rows changed rate')
for k, bv, av in changed[:30]:
    print(f'  {k[1][:55]:<55}  {bv} -> {av}')
