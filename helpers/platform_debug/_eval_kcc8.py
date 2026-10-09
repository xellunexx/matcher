# -*- coding: utf-8 -*-
import sys, io, json, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\ochak\Downloads\kcc2 (8).xlsx', read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
wb.close()
hdr = rows[1]
cols = {n: i for i, n in enumerate(hdr) if n}
i_d = cols['Description']; i_u = cols['Unit']; i_q = cols['Quantity']
i_r = cols['Unit Rate']; i_t = cols['Total']; i_s = cols['Source']
i_c = cols['Confidence']; i_m = cols.get('Metadata JSON')

priced = unpriced = zero = 0
tot = 0.0
rates = []
codes = collections.Counter()
conf = collections.Counter()
srcs = collections.Counter()
no_rate_desc = []
big = []
for r in rows[2:]:
    d = r[i_d]
    if d is None or not str(d).strip():
        continue
    srcs[r[i_s]] += 1
    conf[r[i_c]] += 1
    try:
        v = float(r[i_r]) if r[i_r] is not None else None
    except (TypeError, ValueError):
        v = None
    if v is None or v == 0:
        unpriced += 1
        if v == 0: zero += 1
        no_rate_desc.append(str(d)[:70])
        continue
    priced += 1
    rates.append(v)
    tot += float(r[i_t] or 0)
    if v > 500: big.append((str(d)[:60], v))
    if i_m is not None and r[i_m]:
        try:
            cm = json.loads(r[i_m]).get('cost_match') or {}
            code = str(cm.get('code') or cm.get('ref') or '?')
            codes[code.split('-')[0][:12]] += 1
        except Exception:
            codes['<unparsable>'] += 1

print(f'rows={priced+unpriced} priced={priced} unpriced={unpriced} (zeros={zero}) total={tot:.2f} EUR')
print('sources:', dict(srcs))
print('confidence:', dict(conf))
print('cost_match code prefixes (top 15):')
for k, c in codes.most_common(15): print(f'  {c:>4}  {k}')
print('rates>500:')
for d, v in big[:15]: print(f'  {v:>12,.2f}  {d}')
print('top repeated rates:', collections.Counter(round(x,2) for x in rates).most_common(12))
print(f'--- {len(no_rate_desc)} unpriced (first 25):')
for d in no_rate_desc[:25]: print('  ', d)
