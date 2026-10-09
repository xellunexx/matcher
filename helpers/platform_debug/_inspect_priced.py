import sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\ochak\Music\tendererp\kcc2_priced.xlsx", read_only=True)
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True, max_row=6))
    print(f"\n=== {name} dims={ws.max_row}x{ws.max_column}")
    for r in rows:
        print("   ", [str(c)[:30] if c is not None else "" for c in r][:14])
wb.close()
