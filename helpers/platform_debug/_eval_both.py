# -*- coding: utf-8 -*-
import sys, io, collections
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

def rows_of(p, sheet=None):
    wb = openpyxl.load_workbook(p, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows

# kcc2_priced.xlsx: original KCC layout - desc B, unit C, qty D, price F, total G, conf H, tier I
print('===== price_kcc2 output (kcc2_priced.xlsx)')
rows = rows_of(r'C:\lab\tenderops\kcc2_priced.xlsx')
tiers = collections.Counter()
priced = 0
tot = 0.0
suspicious = []
for r in rows[2:]:
    if len(r) < 9: continue
    desc, unit, qty, rate, ltot, conf, tier = r[1], r[2], r[3], r[5], r[6], r[7], r[8]
    if not desc: continue
    tiers[tier or 'blank'] += 1
    if rate:
        priced += 1
        tot += float(ltot or 0)
        if float(rate) > 200:
            suspicious.append((str(desc)[:60], rate, tier))
print('tiers:', dict(tiers), 'priced:', priced, 'total EUR:', round(tot, 2))
print('rates>200:')
for d, v, t in suspicious[:15]: print(f'  {v:>10} {t}  {d}')

# ERP run kcc2 (6).xlsx: Pos.,Description,Unit,Qty,Unit Rate,Total,...
print('\n===== ERP run (kcc2 (6).xlsx)')
rows = rows_of(r'C:\Users\ochak\Downloads\kcc2 (6).xlsx')
hdr = rows[1]
cols = {n: i for i, n in enumerate(hdr) if n}
i_d, i_u, i_r, i_t, i_s, i_c = cols['Description'], cols['Unit'], cols['Unit Rate'], cols['Total'], cols.get('Source'), cols.get('Confidence')
priced = unpriced = 0
tot = 0.0
zero = 0
rates = []
for r in rows[2:]:
    d = r[i_d]
    if d is None or not str(d).strip(): continue
    try:
        v = float(r[i_r]) if r[i_r] is not None else None
    except (TypeError, ValueError):
        v = None
    if v is None:
        unpriced += 1
    elif v == 0:
        zero += 1; unpriced += 1
    else:
        priced += 1; rates.append(v); tot += float(r[i_t] or 0)
print(f'priced={priced} unpriced={unpriced} (zeros={zero}) total EUR={tot:.2f}')
dist = collections.Counter(round(x, 2) for x in rates)
print('top repeated:', dist.most_common(10))
