# -*- coding: utf-8 -*-
"""Extract priced rows from priced kcc files -> costdb seed JSON -> oe_costs_item
(source='operator_pricelist', OPR3- codes).

Normalization rules (operator spec):
  R1  strip (вед.N) refs entirely
  R2  cut at first "вкл."/"включително" clause -> append ". крайна цена"
      (tag marks the submitted rate was all-inclusive)
  R3  cut "съгласно БДС… / изискванията на ТС" boilerplate tails
  R4  kcc1: use Описание column, never the abbreviated Дейност column
  R5  verbatim case; collapse whitespace/newlines
  R6  merge Характеристика into name only when spec-like (short + discriminator)
"""
import sys, io, json, re, zipfile
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

import openpyxl, xlrd

DIR = r"C:\lab\tenderops\prices"
OUT = DIR + r"\costdb_seed_user_operator_prices3.json"

NAME_RE  = re.compile(r"наимен|описан|видове|смр|дейност|вид\s+стр", re.U)
DESC_RE  = re.compile(r"описан", re.U)          # preferred over Дейност (R4)
UNIT_RE  = re.compile(r"мярка|мяр-\s*ка|мерна|ед\.?\s*м\.?|^\s*ед\s*\.?\s*$", re.U)
CHAR_RE  = re.compile(r"характер", re.U)
QTY_RE   = re.compile(r"к-?во|колич|кол\.|прогнозно", re.U)
UPRICE_RE= re.compile(r"ед\.?\s*цен|единична|цена\s*/", re.U)
PRICE_RE = re.compile(r"цен|ст-?ст|стойност|price|сума", re.U)
EUR_RE   = re.compile(r"€|евро|eur", re.I)
BGN_RE   = re.compile(r"лв|bgn", re.I)

VED_RE   = re.compile(r"\(\s*вед\.?\s*[^)]*\)", re.I|re.U)
VKL_RE   = re.compile(r"[,;.\s]*\b(?:вкл\.?|включително)\b.*$", re.I|re.U|re.S)
BDS_RE   = re.compile(r"[,;.\s]*\bсъгласно\b.*$", re.I|re.U|re.S)
NUMPFX_RE= re.compile(r"^\s*\d+(?:\.\d+)*[.)]?\s+", re.U)
WS_RE    = re.compile(r"\s+", re.U)

# discriminator tokens that make a Характеристика spec-like, not prose
SPEC_RE  = re.compile(r"\d|фракция|клас|деб\.|дiam|dn|pn|ф\s*\d|n\d|w\d|mpa|kn|kw|hz", re.I|re.U)

UNIT_MAP = {"м1": "м", "м.": "м", "м2": "м²", "м3": "м³",
            "m": "м", "m2": "м²", "m3": "м³", "м²": "м²", "м³": "м³",
            "m²": "м²", "m³": "м³", "m\xB2": "м²", "m\xB3": "м³",
            "бр.": "бр.", "бр": "бр.", "бр .": "бр.",
            "кг.": "кг", "kg": "кг", "кг": "кг",
            "т.": "т", "t": "т", "т": "т",
            "ч.ч.": "ч.ч.", "чч": "ч.ч.",
            "компл.": "компл.", "компл": "компл.", "к-т": "компл.",
            "л.": "л", "l": "л"}


def norm_unit(u):
    u = WS_RE.sub(" ", str(u or "").strip().lower().rstrip(","))
    return UNIT_MAP.get(u, u)


NUM_ONLY_RE = re.compile(r"^\d[\d\s.,]*$")


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        t = v.strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
        try:
            return float(t)
        except ValueError:
            return None
    return None


def clean_desc(s):
    """Returns (desc, final_price_tag)."""
    s = WS_RE.sub(" ", str(s or "").replace("\n", " ")).strip()
    s = VED_RE.sub("", s)                       # R1
    final = False
    m = VKL_RE.search(s)
    if m:                                        # R2
        s = s[:m.start()]
        final = True
    s = BDS_RE.sub("", s)                        # R3
    s = NUMPFX_RE.sub("", s)                     # row numbering prefix
    s = s.strip().strip(",;.").strip()
    return s, final


def spec_like(char):
    c = WS_RE.sub(" ", str(char or "").replace("\n", " ")).strip()
    if not c or len(c) > 80:
        return None
    return c if SPEC_RE.search(c) else None


def find_map(rows):
    for i, r in enumerate(rows[:60]):
        cells = [str(c).lower().replace("\n", " ") if c is not None else "" for c in r]
        hdr_text = " ".join(cells)
        name_cands = [j for j, c in enumerate(cells) if NAME_RE.search(c)]
        if not name_cands:
            continue
        # R4: prefer "Описание" over abbreviated "Дейност" when both exist
        name_c = next((j for j in name_cands if DESC_RE.search(cells[j])), name_cands[0])
        unit_c = next((j for j, c in enumerate(cells) if UNIT_RE.search(c)), None)
        char_c = next((j for j, c in enumerate(cells) if CHAR_RE.search(c)), None)
        price_cs = [j for j, c in enumerate(cells) if PRICE_RE.search(c)]
        if not price_cs or unit_c is None:
            continue
        qty_c = next((j for j, c in enumerate(cells) if QTY_RE.search(c)), None)
        up_c = next((j for j, c in enumerate(cells) if UPRICE_RE.search(c)), None)
        if up_c is None:
            cands = [j for j in price_cs if j not in (name_c, unit_c, qty_c)]
            up_c = cands[0] if cands else None
        if up_c is None:
            continue
        tot_c = next((j for j in price_cs if j != up_c), None)
        cur = "EUR" if EUR_RE.search(hdr_text) else ("BGN" if BGN_RE.search(hdr_text) else "EUR")
        return i, dict(name=name_c, unit=unit_c, qty=qty_c, uprice=up_c,
                       total=tot_c, char=char_c), cur
    return None, None, None


def extract_rows(rows, currency, cols, hdr_i):
    out = []
    for r in rows[hdr_i + 1:]:
        g = lambda k: (r[cols[k]] if cols.get(k) is not None and cols[k] < len(r) else None)
        desc, final = clean_desc(g("name"))
        spec = spec_like(g("char"))              # R6
        if spec and spec.lower() not in desc.lower():
            desc = f"{desc}. {spec}".strip(". ")
        unit = norm_unit(str(g("unit") or ""))
        qty, up, tot = num(g("qty")), num(g("uprice")), num(g("total"))
        if len(desc) < 4 or not unit:
            continue
        if NUM_ONLY_RE.match(desc) or NUM_ONLY_RE.match(unit):
            continue  # SAP-code / numeric-artefact rows are not positions
        price, derived = up, False
        if (not price or price <= 0) and tot and qty and qty > 0:
            price, derived = round(tot / qty, 4), True
        if not price or price <= 0:
            continue
        price = round(price, 4)
        if final:
            desc = f"{desc}. крайна цена"
        out.append(dict(desc=desc, unit=unit, amount=price, currency=currency,
                        derived=derived))
    return out


def from_xlsx(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    res = []
    for sn in wb.sheetnames:
        rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
        h, cols, cur = find_map(rows)
        if cols:
            res += extract_rows(rows, cur, cols, h)
    wb.close()
    return res


def from_xls(path):
    wb = xlrd.open_workbook(path)
    res = []
    for sn in wb.sheet_names():
        ws = wb.sheet_by_name(sn)
        rows = [[ws.cell_value(i, j) for j in range(ws.ncols)] for i in range(ws.nrows)]
        h, cols, cur = find_map(rows)
        if cols:
            res += extract_rows(rows, cur, cols, h)
    return res


def from_docx(path):
    z = zipfile.ZipFile(path)
    xml = z.read("word/document.xml").decode("utf-8", "ignore")
    rows = []
    for tr in re.findall(r"<w:tr\b.*?</w:tr>", xml, re.S):
        cells = [" ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", tc)).strip()
                 for tc in re.findall(r"<w:tc\b.*?</w:tc>", tr, re.S)]
        rows.append(cells)
    h, cols, cur = find_map(rows)
    return extract_rows(rows, cur, cols, h) if cols else []


FILES = [
    ("kcc1.xlsx", from_xlsx), ("kcc2.xls", from_xls), ("kcc4.docx", from_docx),
    ("kcc6.xlsx", from_xlsx), ("kcc9.xlsx", from_xlsx), ("kcc11.xlsx", from_xlsx),
    ("kcc13.xlsx", from_xlsx), ("kcc15.xlsx", from_xlsx), ("kcc17.xls", from_xls),
]

if __name__ == "__main__":
    preview = "--preview" in sys.argv
    seed, n = [], 0
    for fname, fn in FILES:
        rows = fn(f"{DIR}\\{fname}")
        for r in rows:
            n += 1
            code = f"OPR3-{n:06d}"
            seed.append({
                "id": code, "code": code, "desc": r["desc"], "name": r["desc"],
                "category": f"operator/{fname.rsplit('.',1)[0]}",
                "categoryLabel": "",
                "unit": r["unit"],
                "money": {"amount": r["amount"], "currency": r["currency"],
                          "vatIncluded": None, "asOf": None},
                "vatStatus": "unknown",
                "origin": {"kind": "operator", "ref": fname,
                           "note": "derived_unit_price" if r["derived"] else ""},
                "validity": {"from": None, "to": None},
                "status": "active",
                "region": "BG", "type": "priced_offer", "approvals": [],
            })
        print(f"{fname}: {len(rows)} priced rows")
        if preview:
            for s in seed[-min(4, len(rows)):]:
                print(f"   {s['code']} | {s['desc'][:95]} | {s['unit']} | "
                      f"{s['money']['amount']} {s['money']['currency']}")
    json.dump(seed, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    from collections import Counter
    print(f"\nseed -> {OUT}: {len(seed)} items")
    print("units:", Counter(s["unit"] for s in seed).most_common(15))
    print("curr:", Counter(s["money"]["currency"] for s in seed))
    print("крайна цена tagged:", sum(1 for s in seed if "крайна цена" in s["desc"]))
