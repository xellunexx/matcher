import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\ochak\Downloads\test.xlsx")
print("sheets:", wb.sheetnames)
ws = wb["BOQ"]
print("dims:", ws.max_row, ws.max_column)
for row in list(ws.iter_rows(max_row=12, values_only=True)):
    print([str(v)[:40] if v is not None else None for v in row])
