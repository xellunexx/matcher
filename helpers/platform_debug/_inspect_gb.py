import sys, io, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl
from decimal import Decimal

wb = openpyxl.load_workbook(r"C:\Users\ochak\Downloads\gb.xlsx")
print("sheets:", wb.sheetnames)
ws = wb[wb.sheetnames[0]]
print("dims:", ws.max_row, ws.max_column)
for row in list(ws.iter_rows(max_row=4, values_only=True)):
    print([str(v)[:60] if v is not None else None for v in row])

# stats
priced = zero = 0
total = Decimal(0)
cur_seen = {}
basis = {}
sources = {}
for row in ws.iter_rows(min_row=3, values_only=True):
    desc, unit, qty, rate, tot, cur = row[1], row[2], row[3], row[4], row[5], row[6]
    if desc is None or unit is None:
        continue
    try:
        r = Decimal(str(rate or 0))
        t = Decimal(str(tot or 0))
    except Exception:
        continue
    if cur: cur_seen[cur] = cur_seen.get(cur, 0) + 1
    if r > 0:
        priced += 1
        total += t
    else:
        zero += 1
    meta = row[13]
    if meta:
        try:
            m = json.loads(meta)
            cm = m.get("cost_match") or {}
            t2 = cm.get("tier") or m.get("price_basis")
            if t2: basis[t2] = basis.get(t2, 0) + 1
            cd = cm.get("code") or ""
            if cd: sources[cd.split("-")[0].split("_")[0]] = sources.get(cd.split("-")[0].split("_")[0], 0) + 1
        except Exception:
            pass
print(f"\npriced={priced} zero={zero} total={total}")
print("currencies:", cur_seen)
print("tiers:", basis)
print("code prefixes:", dict(sorted(sources.items(), key=lambda x:-x[1])[:10]))
